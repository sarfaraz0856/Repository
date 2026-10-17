import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/Repository/excel-summary-tutorial/SUMIF_SUMIFS_Practice_File.xlsx"
HDR_FILL = PatternFill("solid", fgColor="1F4E78"); HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
TOT_FILL = PatternFill("solid", fgColor="C6EFCE"); TOT_FONT = Font(bold=True, color="006100")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="A6A6A6"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

ROWS = [("01-Oct-2026","Rahul Sharma","Laptop",100,20,5000),("01-Oct-2026","Priya Das","Mobile",80,10,6400),
        ("02-Oct-2026","Rahul Sharma","Laptop",50,15,2500),("02-Oct-2026","Aman Khan","Tablet",60,12,7200),
        ("03-Oct-2026","Priya Das","Mobile",40,5,3200),("03-Oct-2026","Rahul Sharma","Tablet",30,8,3600),
        ("04-Oct-2026","Aman Khan","Laptop",70,20,3500),("04-Oct-2026","Priya Das","Mobile",25,6,2000),
        ("05-Oct-2026","Rahul Sharma","Laptop",40,10,2000),("05-Oct-2026","Aman Khan","Tablet",20,5,2400)]
COMBOS = [("Rahul Sharma","Laptop"),("Rahul Sharma","Tablet"),("Priya Das","Mobile"),
          ("Aman Khan","Tablet"),("Aman Khan","Laptop")]
EMPS = ["Rahul Sharma","Priya Das","Aman Khan"]

wb = Workbook()

# ---------- Sheet 1: All Data ----------
ad = wb.active; ad.title = "All Data"
ad.append(["Date","Employee Name","Stock Item","In Qty","Out Qty","Amount"])
for d,*rest in ROWS:
    ad.append([dt.datetime.strptime(d,"%d-%b-%Y")]+rest)
for c in ad[1]:
    c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = Alignment(horizontal="center")
for r in range(2, 12):
    ad.cell(r,1).number_format = "dd-mmm-yyyy"
    ad.cell(r,6).number_format = "#,##0"
    for c in range(1,7): ad.cell(r,c).border = BOX
for col,w in zip("ABCDEF",[14,18,14,10,10,12]): ad.column_dimensions[col].width = w
ad.freeze_panes = "A2"

def summary_sheet(ws, solved):
    ws["A1"] = "EMPLOYEE AND STOCK SUMMARY REPORT"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws["G1"] = "EMPLOYEE-WISE SUMMARY (SUMIF)"
    ws["G1"].font = Font(bold=True, size=14, color="1F4E78")
    for i,h in enumerate(["Employee Name","Stock Item","Total In Qty","Total Out Qty","Total Amount"]):
        ws.cell(2,1+i,h)
    for i,(e,k) in enumerate(COMBOS):
        r = 3+i; ws.cell(r,1,e); ws.cell(r,2,k)
        if solved:
            for c,src in zip("CDE","DEF"):
                ws[f"{c}{r}"] = f"=SUMIFS('All Data'!${src}:${src},'All Data'!$B:$B,$A{r},'All Data'!$C:$C,$B{r})"
    ws["A8"] = "GRAND TOTAL"
    if solved:
        for c in "CDE": ws[f"{c}8"] = f"=SUM({c}3:{c}7)"
    for i,h in enumerate(["Employee Name","Total In Qty","Total Out Qty","Total Amount"]):
        ws.cell(2,7+i,h)
    for i,e in enumerate(EMPS):
        r = 3+i; ws.cell(r,7,e)
        if solved:
            for c,src in zip("HIJ","DEF"):
                ws[f"{c}{r}"] = f"=SUMIF('All Data'!$B:$B,$G{r},'All Data'!${src}:${src})"
    ws["G6"] = "GRAND TOTAL"
    if solved:
        for c in "HIJ": ws[f"{c}6"] = f"=SUM({c}3:{c}5)"
    for rng,total_row in (("A2:E8",8), ("G2:J6",6)):
        for row in ws[rng]:
            for cell in row:
                cell.border = BOX
                if cell.row == 2:
                    cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.alignment = Alignment(horizontal="center")
                elif cell.row == total_row:
                    cell.fill = TOT_FILL; cell.font = TOT_FONT
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.number_format = "#,##0"
    for col,w in zip("ABCDEFGHIJ",[18,14,13,14,14,4,18,13,14,14]): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"

# ---------- Sheet 2: Summary (Practice) ----------
pr = wb.create_sheet("Summary (Practice)"); summary_sheet(pr, solved=False)
pr["A10"] = "👉 PRACTICE: C3:E7 aur H3:J5 mein formula khud lagaiye."
pr["A10"].font = Font(bold=True, size=11, color="7F6000"); pr["A10"].fill = NOTE_FILL
pr["A11"] = "C3  =SUMIFS('All Data'!$D:$D,'All Data'!$B:$B,$A3,'All Data'!$C:$C,$B3)"
pr["A12"] = "H3  =SUMIF('All Data'!$B:$B,$G3,'All Data'!$D:$D)"
pr["A13"] = "Answers check karne ke liye 'Summary (Solution)' sheet dekhiye."
for r in (11,12,13):
    pr[f"A{r}"].font = Font(size=10, color="7F6000", name="Consolas" if r < 13 else "Calibri")
    pr[f"A{r}"].fill = NOTE_FILL

# ---------- Sheet 3: Summary (Solution) ----------
sol = wb.create_sheet("Summary (Solution)"); summary_sheet(sol, solved=True)

wb.save(OUT)
print("saved", OUT)
