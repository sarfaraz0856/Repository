#!/usr/bin/env python3
"""
Build Sales_MIS_Report.xlsx for the Excel 365 MIS tutorial.

Creates a real .xlsx (Open XML) workbook compatible with Microsoft Excel 365,
using openpyxl. Contains a native Excel Table (SalesData) with structured-
reference formulas and a dashboard driven by SUMIFS/COUNTIFS.

Pivot Tables and Slicers are NOT written by this script: openpyxl cannot
author genuine native Excel Pivot Tables or Slicers. The source data and
exact manual steps for those are documented in Read_Me and the guides.
"""
import random
from datetime import date, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ------------------------------------------------------------------ settings
random.seed(20261001)  # deterministic => reproducible, documented data

REGIONS = ["East", "West", "North", "South"]
PRODUCTS = ["Tea A", "Tea B", "Tea C"]
UNIT_PRICE = {"Tea A": 120.0, "Tea B": 95.0, "Tea C": 150.0}  # fictional
STATUS = ["Paid", "Pending"]
SALESPEOPLE = {
    "East": ["Rahul Sharma", "Pooja Nair"],
    "West": ["Amit Patel", "Sneha Rao"],
    "North": ["Vikram Singh", "Neha Gupta"],
    "South": ["Arjun Reddy", "Divya Menon"],
}

N_ROWS = 180  # within the requested 150-200 range
MONTH_START = date(2026, 10, 1)
MONTH_DAYS = 31  # October has 31 days

HEADERS = ["Date", "Invoice No", "Region", "Salesperson", "Product",
           "Qty", "Unit Price", "Sales Amount", "Payment Status"]

# ------------------------------------------------------------------ styling
NAVY = "1F3864"
BLUE = "2E5496"
LIGHT = "D9E1F2"
ACCENT = "C9DAF8"
WHITE = "FFFFFF"
GREY = "7F7F7F"
GREEN = "548235"
AMBER = "BF8F00"

def fill(hex_):
    return PatternFill("solid", fgColor=hex_)

thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
INR = '#,##0.00'           # plain number, 2 dp (keeps file locale-neutral)
INR0 = '#,##0'

# ------------------------------------------------------------------ data gen
def make_rows():
    rows = []
    for i in range(N_ROWS):
        d = MONTH_START + timedelta(days=random.randint(0, MONTH_DAYS - 1))
        region = random.choice(REGIONS)
        sp = random.choice(SALESPEOPLE[region])
        prod = random.choice(PRODUCTS)
        qty = random.randint(5, 120)
        price = UNIT_PRICE[prod]
        status = random.choices(STATUS, weights=[72, 28])[0]
        rows.append([d, region, sp, prod, qty, price, status])
    rows.sort(key=lambda r: r[0])  # chronological
    # assign unique, non-blank invoice numbers AFTER sort
    out = []
    for n, r in enumerate(rows, start=1):
        d, region, sp, prod, qty, price, status = r
        inv = f"INV-2026-{n:04d}"
        out.append([d, inv, region, sp, prod, qty, price, status])
    return out

# ------------------------------------------------------------------ build
wb = openpyxl.Workbook()

# ===== Read_Me =====================================================
ws = wb.active
ws.title = "Read_Me"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 100

def rm(r, text, *, size=11, bold=False, color="000000", top=0):
    c = ws.cell(row=r, column=2, value=text)
    c.font = Font(name="Calibri", size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if top:
        ws.row_dimensions[r].height = top
    return c

rm(2, "SALES MIS REPORT  —  Read Me", size=18, bold=True, color=NAVY)
rm(3, "Reporting period: October 2026   |   Workbook: Sales_MIS_Report.xlsx")
rm(5, "WHAT THIS WORKBOOK CONTAINS", size=13, bold=True, color=BLUE)
readme_lines = [
    "This workbook is a teaching sample for a Sales MIS (Management Information System) report built in Microsoft Excel 365.",
    "",
    "Worksheets:",
    "  1. Read_Me          - this sheet (assumptions, instructions, limitations).",
    "  2. Raw_Data         - the transaction data, stored as an Excel Table named 'SalesData'.",
    "  3. MIS_Dashboard    - KPI cards, Region-wise and Product-wise summaries (all formula-driven).",
    "  4. Pivot_Analysis   - placeholder + exact manual steps to build native Pivot Tables, Charts and Slicers.",
    "",
    "DATA ASSUMPTIONS (IMPORTANT):",
    "  - All data is fictional sample data for October 2026. No real customer or employee information is used.",
    "  - Each row represents exactly ONE unique invoice.",
    "  - Every 'Invoice No' is non-blank and unique (format INV-2026-0001 ... INV-2026-0180).",
    "  - Because one row = one invoice, Total Orders is counted with COUNTIFS on non-blank invoice numbers.",
    "    Do NOT reuse simple row counting as an invoice count if you later add multiple rows per invoice.",
    "",
    "COLUMNS in SalesData:",
    "  Date | Invoice No | Region | Salesperson | Product | Qty | Unit Price | Sales Amount | Payment Status",
    "",
    "KEY FORMULAS:",
    "  - Sales Amount (per row, in the table):   =[@Qty]*[@[Unit Price]]",
    "  - Monthly Total Sales:  =SUMIFS(SalesData[Sales Amount],SalesData[Date],\">=\"&DATE(2026,10,1),SalesData[Date],\"<\"&DATE(2026,11,1))",
    "  - Monthly Total Orders: =COUNTIFS(SalesData[Date],\">=\"&DATE(2026,10,1),SalesData[Date],\"<\"&DATE(2026,11,1),SalesData[Invoice No],\"<>\")",
    "  - Total Units Sold:     =SUMIFS(SalesData[Qty], ... same date criteria ...)",
    "  - Average Sales/Order:  =IFERROR(Total Sales / Total Orders, 0)",
    "  - Pending Sales Amount: =SUMIFS(SalesData[Sales Amount],SalesData[Payment Status],\"Pending\", ... date criteria ...)",
    "",
    "TERMINOLOGY NOTE:",
    "  The pending KPI is labelled 'Pending Sales Amount' (invoiced sales not yet marked Paid).",
    "  It is NOT 'Outstanding Receivables', because this data has no payments-received or balance columns.",
    "",
    "PIVOT TABLES & SLICERS:",
    "  openpyxl (the library that generated this file) cannot create genuine native Excel Pivot Tables or Slicers.",
    "  This script therefore did NOT fabricate fake ones. The 'Pivot_Analysis' sheet and the recording/editing",
    "  guides contain the exact click-by-click steps to create real Pivot Tables, Charts and Slicers in Excel 365.",
    "",
    "HOW TO VERIFY THE NUMBERS:",
    "  Open this file in Microsoft Excel 365. Excel will recalculate all formulas on open.",
    "  Cross-check the dashboard against the QA checklist (scripts/qa_checklist.md).",
]
r = 6
for line in readme_lines:
    rm(r, line, size=11)
    r += 1

# ===== Raw_Data ====================================================
rd = wb.create_sheet("Raw_Data")
rd.sheet_view.showGridLines = False
data = make_rows()

# title band
rd.merge_cells("A1:I1")
t = rd["A1"]
t.value = "RAW DATA  —  SalesData (one row = one unique invoice)"
t.font = Font(size=13, bold=True, color=WHITE)
t.fill = fill(NAVY)
t.alignment = Alignment(vertical="center", indent=1)
rd.row_dimensions[1].height = 24

header_row = 2
for col, h in enumerate(HEADERS, start=1):
    c = rd.cell(row=header_row, column=col, value=h)
    c.font = Font(bold=True, color=WHITE)
    c.fill = fill(BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box

first_data = header_row + 1
for i, row in enumerate(data):
    xr = first_data + i
    d, inv, region, sp, prod, qty, price, status = row
    rd.cell(row=xr, column=1, value=d).number_format = "dd-mmm-yyyy"
    rd.cell(row=xr, column=2, value=inv)
    rd.cell(row=xr, column=3, value=region)
    rd.cell(row=xr, column=4, value=sp)
    rd.cell(row=xr, column=5, value=prod)
    rd.cell(row=xr, column=6, value=qty).number_format = INR0
    rd.cell(row=xr, column=7, value=price).number_format = INR
    # Sales Amount: Excel Table structured-reference formula (col H)
    rd.cell(row=xr, column=8,
            value="=[@Qty]*[@[Unit Price]]").number_format = INR
    rd.cell(row=xr, column=9, value=status)
    for col in range(1, 10):
        rd.cell(row=xr, column=col).border = box

last_data = first_data + len(data) - 1

# column widths
widths = [13, 15, 10, 16, 10, 8, 11, 14, 15]
for col, w in enumerate(widths, start=1):
    rd.column_dimensions[get_column_letter(col)].width = w

# define the Excel Table "SalesData"
ref = f"A{header_row}:I{last_data}"
tbl = Table(displayName="SalesData", ref=ref)
tbl.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2", showRowStripes=True,
    showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
rd.add_table(tbl)
rd.freeze_panes = f"A{first_data}"

# ===== MIS_Dashboard ===============================================
db = wb.create_sheet("MIS_Dashboard")
db.sheet_view.showGridLines = False
for col in range(1, 10):
    db.column_dimensions[get_column_letter(col)].width = 15
db.column_dimensions["A"].width = 3

# title
db.merge_cells("B2:I2")
c = db["B2"]; c.value = "SALES MIS REPORT"
c.font = Font(size=22, bold=True, color=WHITE)
c.fill = fill(NAVY); c.alignment = Alignment(horizontal="center", vertical="center")
db.row_dimensions[2].height = 36
db.merge_cells("B3:I3")
c = db["B3"]; c.value = "Reporting Period:  October 2026"
c.font = Font(size=12, bold=True, color=BLUE)
c.alignment = Alignment(horizontal="center")
db.row_dimensions[3].height = 20

DATE_CRIT = ('SalesData[Date],">="&DATE(2026,10,1),'
             'SalesData[Date],"<"&DATE(2026,11,1)')

kpis = [
    ("Total Sales",            f'=SUMIFS(SalesData[Sales Amount],{DATE_CRIT})', INR),
    ("Total Orders",           f'=COUNTIFS({DATE_CRIT},SalesData[Invoice No],"<>")', INR0),
    ("Total Units Sold",       f'=SUMIFS(SalesData[Qty],{DATE_CRIT})', INR0),
    ("Average Sales / Order",  None, INR),   # filled after we know cells
    ("Pending Sales Amount",   f'=SUMIFS(SalesData[Sales Amount],SalesData[Payment Status],"Pending",{DATE_CRIT})', INR),
]

# KPI cards laid out across B..I, two rows of cards (label row + value row)
# card columns: B-C, D-E, F-G, H-I  (4 per band) -> 5 KPIs => band1 4 cards, band2 1 card
card_cols = [("B", "C"), ("D", "E"), ("F", "G"), ("H", "I")]
start_row = 5
value_cell_map = {}

def place_card(idx_in_band, band_row, label, formula, fmt):
    lc, rc = card_cols[idx_in_band]
    lrow, vrow = band_row, band_row + 1
    db.merge_cells(f"{lc}{lrow}:{rc}{lrow}")
    lcell = db[f"{lc}{lrow}"]
    lcell.value = label
    lcell.font = Font(size=10, bold=True, color=WHITE)
    lcell.fill = fill(BLUE)
    lcell.alignment = Alignment(horizontal="center", vertical="center")
    db.merge_cells(f"{lc}{vrow}:{rc}{vrow}")
    vcell = db[f"{lc}{vrow}"]
    vcell.fill = fill(LIGHT)
    vcell.font = Font(size=16, bold=True, color=NAVY)
    vcell.alignment = Alignment(horizontal="center", vertical="center")
    vcell.number_format = fmt
    if formula:
        vcell.value = formula
    db.row_dimensions[lrow].height = 18
    db.row_dimensions[vrow].height = 30
    return f"{lc}{vrow}"

# band 1: first 4 KPIs
for i in range(4):
    label, formula, fmt = kpis[i]
    cellref = place_card(i, start_row, label, formula, fmt)
    value_cell_map[label] = cellref

# Average Sales/Order references Total Sales & Total Orders cells
tot_sales = value_cell_map["Total Sales"]
tot_orders = value_cell_map["Total Orders"]
avg_ref = value_cell_map["Average Sales / Order"]
db[avg_ref].value = f"=IFERROR({tot_sales}/{tot_orders},0)"

# band 2: 5th KPI (Pending)
label, formula, fmt = kpis[4]
cellref = place_card(0, start_row + 3, label, formula, fmt)
value_cell_map[label] = cellref

# ---- Region-wise summary ----
rstart = start_row + 6  # row 11
db.merge_cells(f"B{rstart}:D{rstart}")
c = db[f"B{rstart}"]; c.value = "REGION-WISE SALES (October 2026)"
c.font = Font(bold=True, color=WHITE); c.fill = fill(BLUE)
c.alignment = Alignment(indent=1, vertical="center")
db.row_dimensions[rstart].height = 20
# headers
db[f"B{rstart+1}"] = "Region"; db[f"C{rstart+1}"] = "Sales Amount"
db[f"D{rstart+1}"] = "% of Total"
for col in ("B", "C", "D"):
    hc = db[f"{col}{rstart+1}"]
    hc.font = Font(bold=True, color=NAVY); hc.fill = fill(ACCENT)
    hc.border = box; hc.alignment = Alignment(horizontal="center")
region_first = rstart + 2
for i, reg in enumerate(REGIONS):
    rr = region_first + i
    db[f"B{rr}"] = reg
    db[f"C{rr}"] = (f'=SUMIFS(SalesData[Sales Amount],SalesData[Region],B{rr},'
                    f'{DATE_CRIT})')
    db[f"C{rr}"].number_format = INR
    db[f"D{rr}"] = f'=IFERROR(C{rr}/{tot_sales},0)'
    db[f"D{rr}"].number_format = "0.0%"
    for col in ("B", "C", "D"):
        db[f"{col}{rr}"].border = box
reg_total = region_first + len(REGIONS)
db[f"B{reg_total}"] = "Total"
db[f"C{reg_total}"] = f"=SUM(C{region_first}:C{region_first+len(REGIONS)-1})"
db[f"C{reg_total}"].number_format = INR
db[f"D{reg_total}"] = f'=IFERROR(C{reg_total}/{tot_sales},0)'
db[f"D{reg_total}"].number_format = "0.0%"
for col in ("B", "C", "D"):
    cc = db[f"{col}{reg_total}"]
    cc.font = Font(bold=True, color=NAVY); cc.fill = fill(LIGHT); cc.border = box

# ---- Product-wise summary ---- (placed to the right)
db.merge_cells(f"F{rstart}:H{rstart}")
c = db[f"F{rstart}"]; c.value = "PRODUCT-WISE SALES (October 2026)"
c.font = Font(bold=True, color=WHITE); c.fill = fill(BLUE)
c.alignment = Alignment(indent=1, vertical="center")
db[f"F{rstart+1}"] = "Product"; db[f"G{rstart+1}"] = "Sales Amount"
db[f"H{rstart+1}"] = "% of Total"
for col in ("F", "G", "H"):
    hc = db[f"{col}{rstart+1}"]
    hc.font = Font(bold=True, color=NAVY); hc.fill = fill(ACCENT)
    hc.border = box; hc.alignment = Alignment(horizontal="center")
prod_first = rstart + 2
for i, prod in enumerate(PRODUCTS):
    rr = prod_first + i
    db[f"F{rr}"] = prod
    db[f"G{rr}"] = (f'=SUMIFS(SalesData[Sales Amount],SalesData[Product],F{rr},'
                    f'{DATE_CRIT})')
    db[f"G{rr}"].number_format = INR
    db[f"H{rr}"] = f'=IFERROR(G{rr}/{tot_sales},0)'
    db[f"H{rr}"].number_format = "0.0%"
    for col in ("F", "G", "H"):
        db[f"{col}{rr}"].border = box
prod_total = prod_first + len(PRODUCTS)
db[f"F{prod_total}"] = "Total"
db[f"G{prod_total}"] = f"=SUM(G{prod_first}:G{prod_first+len(PRODUCTS)-1})"
db[f"G{prod_total}"].number_format = INR
db[f"H{prod_total}"] = f'=IFERROR(G{prod_total}/{tot_sales},0)'
db[f"H{prod_total}"].number_format = "0.0%"
for col in ("F", "G", "H"):
    cc = db[f"{col}{prod_total}"]
    cc.font = Font(bold=True, color=NAVY); cc.fill = fill(LIGHT); cc.border = box

# footer note
note_row = prod_total + 2
db.merge_cells(f"B{note_row}:I{note_row}")
nc = db[f"B{note_row}"]
nc.value = ("All values are calculated live from the SalesData table using "
            "SUMIFS / COUNTIFS. 'Pending Sales Amount' = invoiced sales not yet "
            "marked Paid (not outstanding receivables).")
nc.font = Font(size=9, italic=True, color=GREY)
nc.alignment = Alignment(wrap_text=True, vertical="top")
db.row_dimensions[note_row].height = 28

# ===== Pivot_Analysis ==============================================
pv = wb.create_sheet("Pivot_Analysis")
pv.sheet_view.showGridLines = False
pv.column_dimensions["A"].width = 3
pv.column_dimensions["B"].width = 100
pv.merge_cells("B2:B2")
h = pv["B2"]; h.value = "PIVOT ANALYSIS  —  build these natively in Excel 365"
h.font = Font(size=15, bold=True, color=NAVY)
pv_lines = [
    "",
    "openpyxl cannot create genuine native Pivot Tables or Slicers, so none were fabricated here.",
    "Follow these exact steps in Microsoft Excel 365 (the source table 'SalesData' is ready on Raw_Data).",
    "",
    "PIVOT TABLE 1  —  Region-wise Sales",
    "  1. Click any cell inside the SalesData table (Raw_Data sheet).",
    "  2. Insert > PivotTable > From Table/Range. Table/Range shows 'SalesData'.",
    "  3. Place it on 'New Worksheet' (or on this sheet at cell B20).",
    "  4. Drag 'Region' to Rows.",
    "  5. Drag 'Sales Amount' to Values. Click it > Value Field Settings > Sum (NOT Count).",
    "  6. Drag 'Date' to Filters. Open the filter and restrict to October 2026",
    "     (or group the Date field by Months and pick Oct).",
    "",
    "PIVOT TABLE 2  —  Product-wise Sales",
    "  1. Insert another PivotTable from SalesData.",
    "  2. Drag 'Product' to Rows.",
    "  3. Drag 'Sales Amount' to Values as Sum (verify it is not Count).",
    "  4. Drag 'Date' to Filters and set October 2026.",
    "",
    "CHARTS",
    "  - Region-wise Sales Performance: select PivotTable 1 > PivotChart > Clustered Column.",
    "  - Product-wise Sales Analysis:   select PivotTable 2 > PivotChart > Pie or Bar.",
    "  - Add data labels: right-click series > Add Data Labels.",
    "",
    "SLICERS",
    "  1. Click a PivotTable > PivotTable Analyze > Insert Slicer.",
    "  2. Tick 'Region' and 'Product' to create two slicers.",
    "  3. To connect a slicer to BOTH pivots (they must share the same data source):",
    "     right-click the slicer > Report Connections (or Filter Connections) >",
    "     tick every PivotTable you want it to control.",
    "  Note: Slicers only cross-filter pivots built from the SAME source (the SalesData table / one cache).",
    "",
    "VERIFY: Sum of each pivot's 'Sales Amount' for October should match the MIS_Dashboard KPI 'Total Sales'.",
    "REFRESH: after editing Raw_Data, PivotTable Analyze > Refresh All (Ctrl+Alt+F5) to update pivots.",
]
r = 3
for line in pv_lines:
    cc = pv.cell(row=r, column=2, value=line)
    bold = line.strip().startswith(("PIVOT TABLE", "CHARTS", "SLICERS", "VERIFY", "REFRESH"))
    cc.font = Font(size=11, bold=bold, color=(BLUE if bold else "000000"))
    cc.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

# ------------------------------------------------------------------ save
wb.properties.title = "Sales MIS Report - October 2026"
wb.properties.creator = "Excel MIS Tutorial"
out = "/home/user/Repository/Excel_MIS_Tutorial/workbook/Sales_MIS_Report.xlsx"
wb.save(out)
print("Saved:", out)
print("Rows:", len(data), "| first invoice:", data[0][1], "| last:", data[-1][1])
