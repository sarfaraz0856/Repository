#!/usr/bin/env python3
"""
Practice files for viewers of the Sales MIS Report video.

  practice/Sales_MIS_Practice.xlsx - same workbook as the video, with the work removed:
      * Raw_Data: the same 180 rows, NOT yet an Excel Table, Sales Amount (H) empty
      * MIS_Dashboard: same layout and formatting, every result cell empty
      * Start_Here: the tasks in video order, and a self-check that marks each
        answer Correct / Check again
  practice/Sales_MIS_Solution.xlsx - the finished workbook from the video

Cells the viewer fills are light yellow (legend on Start_Here). Expected answers
on Start_Here are the workbook's own validated results (scripts/qa_checklist.md).

Usage: python3 scripts/make_practice_file.py
"""
import os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = f"{ROOT}/workbook/Sales_MIS_Report.xlsx"
OUTDIR = f"{ROOT}/practice"
os.makedirs(OUTDIR, exist_ok=True)

FONT = "Calibri"                      # same font as the video workbook
NAVY, BLUE, GREEN, GREY = "1F3864", "2E5496", "107C41", "595959"
FILL_IN = PatternFill("solid", fgColor="FFF2CC")   # light yellow = viewer fills this
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

# Dashboard result cells (top-left of merges) and their number formats
KPI_CELLS = ["B6", "D6", "F6", "H6", "B9"]
REGION_CELLS = [f"{c}{r}" for r in range(13, 18) for c in ("C", "D")]
PRODUCT_CELLS = [f"{c}{r}" for r in range(13, 17) for c in ("G", "H")]
LAST = 182                            # last data row in Raw_Data

# What the finished workbook shows (validated in scripts/qa_checklist.md)
CHECKS = [
    # label, where, formula pulling the viewer's value, expected
    ("Sales Amount, first row (55 × 120)", "Raw_Data!H3", "Raw_Data!H3", 6600),
    ("Sales Amount, whole column (sum)", "Raw_Data!H3:H182", f"SUM(Raw_Data!H3:H{LAST})", 1428190),
    ("Total Sales", "MIS_Dashboard!B6", "MIS_Dashboard!B6", 1428190),
    ("Total Orders", "MIS_Dashboard!D6", "MIS_Dashboard!D6", 180),
    ("Total Units Sold", "MIS_Dashboard!F6", "MIS_Dashboard!F6", 11834),
    ("Average Sales / Order", "MIS_Dashboard!H6", "MIS_Dashboard!H6", 7934.39),
    ("Pending Sales Amount", "MIS_Dashboard!B9", "MIS_Dashboard!B9", 372985),
    ("Region: East", "MIS_Dashboard!C13", "MIS_Dashboard!C13", 351030),
    ("Region: West", "MIS_Dashboard!C14", "MIS_Dashboard!C14", 383675),
    ("Region: North", "MIS_Dashboard!C15", "MIS_Dashboard!C15", 355130),
    ("Region: South", "MIS_Dashboard!C16", "MIS_Dashboard!C16", 338355),
    ("Region total", "MIS_Dashboard!C17", "MIS_Dashboard!C17", 1428190),
    ("Product: Tea A", "MIS_Dashboard!G13", "MIS_Dashboard!G13", 385320),
    ("Product: Tea B", "MIS_Dashboard!G14", "MIS_Dashboard!G14", 432820),
    ("Product: Tea C", "MIS_Dashboard!G15", "MIS_Dashboard!G15", 610050),
    ("Product total", "MIS_Dashboard!G16", "MIS_Dashboard!G16", 1428190),
]

TASKS = [
    ("0:46", "Look at the data",
     "Raw_Data has 180 sales rows for October 2026. One row = one unique invoice."),
    ("1:55", "Make it an Excel Table",
     "Click any cell in A3:I182, press Ctrl + T, keep 'My table has headers' ticked, OK. "
     "Then Table Design → Table Name: SalesData → Enter."),
    ("2:42", "Sales Amount (column H)",
     "In H3 type  =[@Qty]*[@[Unit Price]]  and press Enter. The Table fills all 180 rows. "
     "Example: H3 must show 6,600.00 (55 × 120)."),
    ("3:36", "KPI cards (MIS_Dashboard)",
     "B6 Total Sales: SUMIFS(SalesData[Sales Amount], SalesData[Date], \">=\"&DATE(2026,10,1), "
     "SalesData[Date], \"<\"&DATE(2026,11,1)).  D6 Total Orders: COUNTIFS on the same dates with "
     "SalesData[Invoice No], \"<>\".  F6 Units: SUMIFS on SalesData[Qty].  H6 Average: IFERROR(B6/D6,0).  "
     "B9 Pending: SUMIFS with SalesData[Payment Status], \"Pending\"."),
    ("5:17", "Region-wise and product-wise",
     "C13: SUMIFS(SalesData[Sales Amount], SalesData[Region], B13, + the two date conditions), copy to C16. "
     "C17: SUM(C13:C16). D13: IFERROR(C13/$B$6,0), copy to D17. Same idea for Product in G13:G16 and H13:H16."),
    ("6:19", "Pivot Table",
     "Click inside SalesData → Insert → PivotTable → From Table/Range → New Worksheet. "
     "Region → Rows, Sales Amount → Values (Sum, not Count), Date → Filters. Make a second one with Product."),
    ("7:24", "Pivot Charts",
     "Region pivot → PivotChart → Clustered Column; Product pivot → Pie. Right-click → Add Data Labels."),
    ("8:02", "Slicers",
     "PivotTable Analyze → Insert Slicer → Region and Product. Right-click a slicer → Report Connections → "
     "tick both PivotTables."),
    ("9:24", "Refresh and validate",
     "Data → Refresh All (Ctrl + Alt + F5). Region total = Product total = Pivot Grand Total = Total Sales."),
]


def f(size=11, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def build_practice():
    wb = openpyxl.load_workbook(SRC)

    # ---- Raw_Data: plain range (not a Table yet), Sales Amount empty
    rd = wb["Raw_Data"]
    for name in list(rd.tables.keys()):
        del rd.tables[name]
    for r in range(3, LAST + 1):
        c = rd.cell(row=r, column=8)
        c.value = None
        c.fill = FILL_IN

    # ---- MIS_Dashboard: same layout, results removed
    db = wb["MIS_Dashboard"]
    for ref in KPI_CELLS + REGION_CELLS + PRODUCT_CELLS:
        db[ref].value = None
        db[ref].fill = FILL_IN
    db["B18"].value = ("Fill the yellow cells with formulas (see Start_Here). Never type the numbers by hand: "
                       "the self-check on Start_Here turns green when your answers match the video.")

    # ---- Read_Me (it lists the answer formulas) -> replaced by Start_Here
    del wb["Read_Me"]
    sh = wb.create_sheet("Start_Here", 0)
    build_start_here(sh)
    wb.active = 0

    tidy_pivot_sheet(wb, practice=True)
    fit_to_page_width(sh, landscape=False)
    fit_to_page_width(db)
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "Sales MIS Report - Practice File"
    out = f"{OUTDIR}/Sales_MIS_Practice.xlsx"
    wb.save(out)
    return out


def build_start_here(ws):
    ws.sheet_view.showGridLines = False
    widths = {"A": 3, "B": 9, "C": 30, "D": 22, "E": 16, "F": 16, "G": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("B2:G2")
    ws["B2"] = "SALES MIS REPORT  —  Practice File"
    ws["B2"].font = f(20, True, "FFFFFF")
    ws["B2"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 36
    ws.merge_cells("B3:G3")
    ws["B3"] = "Build the same report as in the video, step by step. Data: 180 fictional sales rows, October 2026."
    ws["B3"].font = f(11, False, BLUE, True)
    ws["B3"].alignment = Alignment(horizontal="center")

    # legend
    ws["B5"] = "Legend"
    ws["B5"].font = f(12, True, GREEN)
    ws["B6"].fill = FILL_IN
    ws["B6"].border = BOX
    ws.merge_cells("C6:G6")
    ws["C6"] = "Light-yellow cell = you fill this with a formula (Raw_Data column H, MIS_Dashboard results)."
    ws["C6"].font = f()
    ws.merge_cells("C7:G7")
    ws["C7"] = "Everything else is already done. The finished answer is in Sales_MIS_Solution.xlsx."
    ws["C7"].font = f(10, False, GREY)

    # tasks
    r = 9
    ws[f"B{r}"] = "Tasks (video time → step)"
    ws[f"B{r}"].font = f(12, True, GREEN)
    r += 1
    for col, head in zip("BCD", ("Video", "Step", "What to do")):
        c = ws[f"{col}{r}"]
        c.value = head
        c.font = f(11, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.border = BOX
    ws.merge_cells(f"D{r}:G{r}")
    for t, step, what in TASKS:
        r += 1
        ws[f"B{r}"] = t
        ws[f"C{r}"] = step
        ws.merge_cells(f"D{r}:G{r}")
        ws[f"D{r}"] = what
        for col in "BCD":
            ws[f"{col}{r}"].font = f(10.5, col == "C")
            ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws[f"{col}{r}"].border = BOX
        for col in "EFG":
            ws[f"{col}{r}"].border = BOX
        ws.row_dimensions[r].height = max(30, 15.5 * (len(what) // 52 + 1))

    # self-check
    r += 2
    ws[f"B{r}"] = "Self-check (updates automatically)"
    ws[f"B{r}"].font = f(12, True, GREEN)
    r += 1
    heads = ("#", "Check", "Cell", "Expected", "Your result", "Status")
    for col, head in zip("BCDEFG", heads):
        c = ws[f"{col}{r}"]
        c.value = head
        c.font = f(11, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
    first = r + 1
    for i, (label, where, pull, expected) in enumerate(CHECKS, 1):
        r += 1
        ws[f"B{r}"] = i
        ws[f"C{r}"] = label
        ws[f"D{r}"] = where
        ws[f"E{r}"] = expected
        if ":" in pull.split("!")[-1]:              # a range: blank until column H has numbers
            rng = pull[pull.index("(") + 1:-1]
            ws[f"F{r}"] = f'=IF(COUNT({rng})=0,"",{pull})'
        else:
            ws[f"F{r}"] = f'=IF(ISBLANK({pull}),"",{pull})'
        ws[f"G{r}"] = f'=IF(F{r}="","",IF(ABS(F{r}-E{r})<0.01,"✔ Correct","✘ Check again"))'
        fmt = "#,##0" if float(expected).is_integer() else "#,##0.00"
        for col in "BCDEFG":
            c = ws[f"{col}{r}"]
            c.font = f(10.5)
            c.border = BOX
        ws[f"B{r}"].alignment = Alignment(horizontal="center")
        ws[f"D{r}"].font = f(9.5, False, GREY)
        ws[f"E{r}"].number_format = fmt
        ws[f"F{r}"].number_format = fmt
        ws[f"G{r}"].alignment = Alignment(horizontal="center")
        ws[f"G{r}"].font = f(10.5, True)
    last = r
    r += 1
    ws.merge_cells(f"B{r}:F{r}")
    ws[f"B{r}"] = "Score"
    ws[f"B{r}"].font = f(11, True)
    ws[f"B{r}"].alignment = Alignment(horizontal="right")
    ws[f"G{r}"] = f'=COUNTIF(G{first}:G{last},"✔ Correct")&" / {len(CHECKS)}"'
    ws[f"G{r}"].font = f(12, True, GREEN)
    ws[f"G{r}"].alignment = Alignment(horizontal="center")
    r += 1
    ws.merge_cells(f"B{r}:G{r}")
    ws[f"B{r}"] = ("Expected values = the results of the finished workbook shown in the video "
                   "(Average Sales / Order is checked to 2 decimals: 1,428,190 ÷ 180 = 7,934.39).")
    ws[f"B{r}"].font = f(9.5, False, GREY, True)
    ws[f"B{r}"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r].height = 28

    # status colours
    from openpyxl.formatting.rule import FormulaRule
    ok = PatternFill("solid", fgColor="C6EFCE")
    bad = PatternFill("solid", fgColor="FFC7CE")
    ws.conditional_formatting.add(f"G{first}:G{last}",
                                  FormulaRule(formula=[f'LEFT(G{first},1)="✔"'], fill=ok, font=Font(color="006100", bold=True)))
    ws.conditional_formatting.add(f"G{first}:G{last}",
                                  FormulaRule(formula=[f'LEFT(G{first},1)="✘"'], fill=bad, font=Font(color="9C0006", bold=True)))
    ws.freeze_panes = "A4"


def fit_to_page_width(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def tidy_pivot_sheet(wb, practice):
    """Replace the build-tool note on Pivot_Analysis with a line for viewers; Calibri throughout."""
    ws = wb["Pivot_Analysis"]
    assert "openpyxl" in str(ws["B4"].value)
    if practice:
        ws["B4"] = ("Do Tasks 1-5 on Start_Here first: Pivot Tables need the data as an Excel Table "
                    "named SalesData (Task 2).")
        ws["B5"] = "Then follow these steps in Microsoft Excel 365:"
    else:
        ws["B4"] = "In this file the data is already the Excel Table 'SalesData'."
        ws["B5"] = ("Pivot Tables, charts and slicers are built in Excel itself (they are not pre-built in "
                    "this file). Follow these steps in Microsoft Excel 365:")
    for row in ws.iter_rows(min_col=2, max_col=2):
        c = row[0]
        if c.value is not None:
            c.font = Font(name=FONT, size=c.font.sz or 11, bold=c.font.b, color=c.font.color)


def build_solution():
    wb = openpyxl.load_workbook(SRC)
    tidy_pivot_sheet(wb, practice=False)
    rm = wb["Read_Me"]                          # build-tool notes -> lines for viewers
    assert "openpyxl" in str(rm["B37"].value) and "qa_checklist" in str(rm["B43"].value)
    rm["B37"] = "  Pivot Tables, charts and slicers are not pre-built in this file: build them in Excel 365"
    rm["B38"] = "  with the steps on the Pivot_Analysis sheet (the data is already the Excel Table 'SalesData')."
    rm["B39"] = None
    rm["B43"] = "  The dashboard should show the same numbers as the video (Total Sales 1,428,190.00)."
    wb.calculation.fullCalcOnLoad = True      # Excel computes every formula when the file opens
    wb.properties.title = "Sales MIS Report - Solution"
    out = f"{OUTDIR}/Sales_MIS_Solution.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build_practice())
    print(build_solution())
