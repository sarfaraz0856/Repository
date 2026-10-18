#!/usr/bin/env python3
"""
Independent validation of Sales_MIS_Report.xlsx.

Part A: structural checks (sheets, columns, table, uniqueness) via openpyxl.
Part B: recompute every KPI straight from the raw data in Python and print the
        reconciliation, so the numbers do not depend on any formula engine.
Part C (optional): if a LibreOffice-recalculated copy is present, read the
        cached formula results and compare them to Part B.
"""
import sys
from datetime import date
import openpyxl

WB = "/home/user/Repository/Excel_MIS_Tutorial/workbook/Sales_MIS_Report.xlsx"
EXPECTED_SHEETS = ["Read_Me", "Raw_Data", "MIS_Dashboard", "Pivot_Analysis"]
EXPECTED_COLS = ["Date", "Invoice No", "Region", "Salesperson", "Product",
                 "Qty", "Unit Price", "Sales Amount", "Payment Status"]
UNIT_PRICE = {"Tea A": 120.0, "Tea B": 95.0, "Tea C": 150.0}

fails, checks = [], []
def ok(name, cond, detail=""):
    checks.append((cond, name, detail))
    if not cond:
        fails.append(name)

# ---------- Part A: structure ----------
wb = openpyxl.load_workbook(WB)
ok("Worksheet names", wb.sheetnames == EXPECTED_SHEETS,
   f"found {wb.sheetnames}")

rd = wb["Raw_Data"]
hdr = [rd.cell(row=2, column=c).value for c in range(1, 10)]
ok("Required columns present", hdr == EXPECTED_COLS, f"found {hdr}")

ok("SalesData table exists", "SalesData" in rd.tables,
   f"tables={list(rd.tables)}")

# read raw rows (Qty, Unit Price are literal; Sales Amount is a formula string)
rows = []
r = 3
while rd.cell(row=r, column=2).value not in (None, ""):
    d = rd.cell(row=r, column=1).value
    inv = rd.cell(row=r, column=2).value
    region = rd.cell(row=r, column=3).value
    prod = rd.cell(row=r, column=5).value
    qty = rd.cell(row=r, column=6).value
    price = rd.cell(row=r, column=7).value
    status = rd.cell(row=r, column=9).value
    rows.append((d, inv, region, prod, qty, price, status))
    r += 1

n = len(rows)
ok("Row count in 150-200 range", 150 <= n <= 200, f"n={n}")

invs = [x[1] for x in rows]
ok("All invoice numbers non-blank", all(bool(str(i).strip()) for i in invs))
ok("All invoice numbers unique", len(set(invs)) == len(invs),
   f"{len(invs)-len(set(invs))} duplicate(s)")

ok("All Qty numeric", all(isinstance(x[4], (int, float)) for x in rows))
ok("All Unit Price numeric", all(isinstance(x[5], (int, float)) for x in rows))
ok("All dates are date objects", all(isinstance(x[0], (date,)) for x in rows))
ok("Unit prices match product map",
   all(abs(x[5] - UNIT_PRICE.get(x[3], -1)) < 1e-9 for x in rows))

def in_oct(d):
    return date(2026, 10, 1) <= (d.date() if hasattr(d, "date") else d) < date(2026, 11, 1)

ok("All dates fall in October 2026", all(in_oct(x[0]) for x in rows))

# ---------- Part B: independent recomputation ----------
sales = {inv: round(q * p, 2) for (_, inv, _, _, q, p, _) in rows}
oct_rows = [x for x in rows if in_oct(x[0])]

total_sales   = round(sum(sales[x[1]] for x in oct_rows), 2)
total_orders  = len({x[1] for x in oct_rows})           # unique invoices
total_units   = sum(x[4] for x in oct_rows)
avg_per_order = round(total_sales / total_orders, 2) if total_orders else 0
pending       = round(sum(sales[x[1]] for x in oct_rows if x[6] == "Pending"), 2)

region_tot = {}
for x in oct_rows:
    region_tot[x[2]] = round(region_tot.get(x[2], 0) + sales[x[1]], 2)
product_tot = {}
for x in oct_rows:
    product_tot[x[3]] = round(product_tot.get(x[3], 0) + sales[x[1]], 2)

ok("Total Orders == unique invoices in Oct", total_orders == len(oct_rows),
   f"orders={total_orders}, oct_rows={len(oct_rows)} (equal because 1 row = 1 invoice)")
ok("Region totals reconcile to Total Sales",
   abs(sum(region_tot.values()) - total_sales) < 0.01,
   f"sum(regions)={sum(region_tot.values()):.2f} vs {total_sales:.2f}")
ok("Product totals reconcile to Total Sales",
   abs(sum(product_tot.values()) - total_sales) < 0.01,
   f"sum(products)={sum(product_tot.values()):.2f} vs {total_sales:.2f}")

# ---------- Part C: compare to LibreOffice-recalculated cache (optional) ----------
import os
RECALC = "/tmp/claude-0/-home-user-Repository/33d85502-7cf9-5098-8064-f58aa938ec5b/scratchpad/recalc/Sales_MIS_Report.xlsx"
recalc_report = []
if os.path.exists(RECALC):
    wb2 = openpyxl.load_workbook(RECALC, data_only=True)
    db = wb2["MIS_Dashboard"]
    # locate KPI value cells by scanning labels
    def find_value_below(label):
        for row in db.iter_rows():
            for cell in row:
                if cell.value == label:
                    return db.cell(row=cell.row + 1, column=cell.column).value
        return None
    pairs = [
        ("Total Sales", total_sales),
        ("Total Orders", total_orders),
        ("Total Units Sold", total_units),
        ("Average Sales / Order", avg_per_order),
        ("Pending Sales Amount", pending),
    ]
    for label, expected in pairs:
        got = find_value_below(label)
        match = got is not None and abs(float(got) - float(expected)) < 0.5
        ok(f"[LibreOffice recalc] {label}", match, f"excel={got} expected={expected}")
        recalc_report.append((label, got, expected))
else:
    recalc_report = None

# ---------- report ----------
print("=" * 64)
print("STRUCTURAL + DATA VALIDATION")
print("=" * 64)
for cond, name, detail in checks:
    print(f"  [{'PASS' if cond else 'FAIL'}] {name}" + (f"  ({detail})" if detail else ""))

print("\n" + "=" * 64)
print("INDEPENDENTLY COMPUTED OCTOBER 2026 KPIs (from raw data)")
print("=" * 64)
print(f"  Total Sales ............. {total_sales:,.2f}")
print(f"  Total Orders ............ {total_orders}")
print(f"  Total Units Sold ........ {total_units:,}")
print(f"  Average Sales / Order ... {avg_per_order:,.2f}")
print(f"  Pending Sales Amount .... {pending:,.2f}")
print("  Region-wise:")
for k in ["East", "West", "North", "South"]:
    print(f"      {k:6} {region_tot.get(k,0):>14,.2f}")
print("  Product-wise:")
for k in ["Tea A", "Tea B", "Tea C"]:
    print(f"      {k:6} {product_tot.get(k,0):>14,.2f}")

if recalc_report is None:
    print("\n[Part C] No LibreOffice-recalculated copy found yet (run recalc step).")

print("\n" + "=" * 64)
print(f"RESULT: {len(checks)-len(fails)}/{len(checks)} checks passed")
if fails:
    print("FAILED:", ", ".join(fails))
    sys.exit(1)
print("ALL CHECKS PASSED")
