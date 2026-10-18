#!/usr/bin/env python3
"""
Build the animated Excel 365 Sales MIS tutorial: 1920x1080, 30 fps, MP4 with audio.

HONEST SCOPE
- ANIMATED RECREATION of the Excel 365 interface (HTML/CSS -> headless Chromium
  -> FFmpeg), not a live Microsoft Excel screen recording. The title bar of
  every frame says "Animated recreation - AI voice".
- Worksheets are rendered from workbook/Sales_MIS_Report.xlsx itself (layout,
  styles, number formats, real formulas in the formula bar). Values come from
  a LibreOffice-recalculated copy and match scripts/qa_checklist.md.
- Pivot/chart/slicer numbers are aggregated here from the same workbook data.
- Voice: AI text-to-speech (Kokoro-82M, Hindi voice, offline), not a human.
- SFX and background music are synthesized locally (scripts/audio_fx.py).

Usage: python3 scripts/build_video.py
"""
import hashlib, html as _html, json, os, re, subprocess, sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import numpy as np
import openpyxl
import soundfile as sf
import excel_ui as ui
import audio_fx as A
import pronunciation

ROOT = os.path.dirname(HERE)
SCR = "/tmp/claude-0/-home-user-Repository/33d85502-7cf9-5098-8064-f58aa938ec5b/scratchpad"
WORK = f"{SCR}/v3"
TTS_MODELS = f"{SCR}/tts"
RECALC = f"{SCR}/recalc/Sales_MIS_Report.xlsx"   # LibreOffice-recalculated copy (values only)
for d in ("html", "png", "clips", "audio", "tts", "assets"):
    os.makedirs(f"{WORK}/{d}", exist_ok=True)
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
FPS = 30
RENDER_H = 1180          # headless Chromium under-paints a 1080 viewport; render taller, crop
SPEED = 1.15         # whole video plays 1.15x faster than the original 0.92-speed build
LEAD, TAIL = 0.45 / SPEED, 0.85 / SPEED
VOICE = os.environ.get("MIS_VOICE", "hm_psi")
VOICE_SPEED = 1.135   # calibrated: narration = 1/1.15 of the 0.92-speed build (Kokoro speed is not linear)
OUT = f"{ROOT}/output/mis_tutorial_animated.mp4"
SRT = f"{ROOT}/subtitles/subtitles.srt"

# ================================================================== workbook
WBS = openpyxl.load_workbook(f"{ROOT}/workbook/Sales_MIS_Report.xlsx")
WBV = openpyxl.load_workbook(RECALC, data_only=True)
RAW = ui.XlsxSource(WBS["Raw_Data"], WBV["Raw_Data"])
DSH = ui.XlsxSource(WBS["MIS_Dashboard"], WBV["MIS_Dashboard"])
RDM = ui.XlsxSource(WBS["Read_Me"], WBV["Read_Me"])

rows = []
r = 3
while WBV["Raw_Data"].cell(r, 2).value:
    c = lambda k: WBV["Raw_Data"].cell(r, k).value
    rows.append(dict(region=c(3), product=c(5), qty=c(6), amt=round(c(6) * c(7), 2)))
    r += 1
LAST = r - 1                                           # 182
assert len(rows) == 180

def agg(key, filt=None):
    d = defaultdict(float)
    for x in rows:
        if filt is None or filt(x):
            d[x[key]] += x["amt"]
    return {k: round(v, 2) for k, v in d.items()}

REG = agg("region")
PRO = agg("product")
WEST_PRO = agg("product", lambda x: x["region"] == "West")
TOTAL = round(sum(x["amt"] for x in rows), 2)
assert abs(TOTAL - 1428190) < 0.01 and abs(sum(REG.values()) - TOTAL) < 0.01
REG_ORDER = sorted(REG)                       # PivotTable sorts row labels A-Z
PRO_ORDER = sorted(PRO)

SHEETS0 = ["Read_Me", "Raw_Data", "MIS_Dashboard", "Pivot_Analysis"]
SHEETS1 = ["Read_Me", "Sheet1", "Raw_Data", "MIS_Dashboard", "Pivot_Analysis"]
Z_RAW, Z_DSH, Z_RDM, Z_PIV = 1.25, 1.45, 1.15, 1.3
TABLE = {"ref": f"A2:I{LAST}", "banded": True, "arrows": True}
H_BLANK = {(rr, 8) for rr in range(3, LAST + 1)}
NUMFMT_NAME = {"#,##0.00": "Number", "#,##0": "Number", "0.0%": "Percentage", "General": "General"}

# ================================================================== frames
_views = {}
def view(src, key, **kw):
    k = (key, tuple(sorted(kw.items())))
    if k not in _views:
        _views[k] = ui.SheetView(src, **kw)
    return _views[k]


class Frame:
    def __init__(self, v, vkw, pkw, overlay=""):
        self.v, self.vkw, self.pkw, self.overlay = v, vkw, pkw, overlay

    def html(self, caption="", badge=""):
        return ui.page(view_html=self.v.html(**self.vkw), overlay=self.overlay,
                       caption=caption, badge=badge, **self.pkw)


def _common(sheet, sheets, tab, ctx, rib, rib_hl, rib_kw, nb, fx, edit, mode, info, zoom):
    return dict(sheet_names=sheets, sheet=sheet, tab=tab, context=ctx, rib=rib, rib_hl=rib_hl,
                rib_kw=rib_kw or {}, namebox=nb, formula=fx, editing=edit is not None,
                mode=mode or ("Enter" if edit is not None else "Ready"), status_info=info, zoom=zoom)


def raw(sel="A3", rng=None, ants=None, table=True, tab="Home", ctx=None, rib=None, rib_hl=None,
        rib_kw=None, edit=None, blank=None, hls=(), top=1, dlg="", mode=None, info="",
        sheets=SHEETS0, nb=None, fx=None):
    v = view(RAW, "raw", zoom=Z_RAW, frozen_rows=2, top_row=top)
    if fx is None:
        if edit is not None:
            fx = edit
        elif sel:
            rr, cc = ui.parse_ref(sel)
            fx = "" if (blank and (rr, cc) in blank) else RAW.formula(rr, cc)
        else:
            fx = ""
    vkw = dict(sel=sel, rng=rng, ants=ants, table=TABLE if table else None, edit=edit,
               blank=blank, hls=tuple(hls), freeze_line=True)
    return Frame(v, vkw, _common("Raw_Data", sheets, tab, ctx, rib, rib_hl, rib_kw,
                                 nb or sel or "A1", fx, edit, mode, info, int(Z_RAW * 100)), dlg)


def dash(sel="A1", rng=None, hls=(), grid=False, tab="Home", ctx=None, rib=None, rib_hl=None,
         rib_kw=None, edit=None, blank=None, dlg="", sheets=SHEETS0, nb=None, fx=None, mode=None, info=""):
    v = view(DSH, "dash", zoom=Z_DSH)
    rr, cc = ui.parse_ref(sel)
    if fx is None:
        fx = edit if edit is not None else ("" if (blank and (rr, cc) in blank) else DSH.formula(rr, cc))
    if rib_kw is None and (rib or tab) == "Home":
        nf = WBS["MIS_Dashboard"].cell(rr, cc).number_format
        rib_kw = {"numfmt": NUMFMT_NAME.get(nf, "Custom")}
    vkw = dict(sel=sel, rng=rng, gridlines=grid, edit=edit, blank=blank, hls=tuple(hls))
    return Frame(v, vkw, _common("MIS_Dashboard", sheets, tab, ctx, rib, rib_hl, rib_kw,
                                 nb or sel, fx, edit, mode, info, int(Z_DSH * 100)), dlg)


def readme(sel="B14", hls=(), sheets=SHEETS0):
    v = view(RDM, "readme", zoom=Z_RDM)
    rr, cc = ui.parse_ref(sel)
    return Frame(v, dict(sel=sel, gridlines=False, hls=tuple(hls)),
                 _common("Read_Me", sheets, "Home", None, None, None, None, sel,
                         RDM.formula(rr, cc), None, None, "", int(Z_RDM * 100)))


# ---------------------------------------------------------------- pivot sheet
HDR = dict(bold=True, fill=ui.BAND, border={"bottom": ui.ACCENT1})
GT = dict(bold=True, border={"top": ui.ACCENT1})


def pivot_source(stage, west=False):
    cells = {}
    def put(r, c, text, num=False, **st):
        d = dict(text=str(text), num=num)
        d.update(st)
        cells[(r, c)] = d
    def block(col, field, data, order, with_values=True, filt=True):
        if filt:
            put(1, col, "Date", bold=True)
            put(1, col + 1, "(All)", arrow=True)
        put(3, col, "Row Labels", arrow=True, **HDR)
        if with_values:
            put(3, col + 1, "Sum of Sales Amount", **HDR)
        rr = 4
        for k in order:
            put(rr, col, k)
            if with_values:
                put(rr, col + 1, ui.fmt(data[k], "General"), num=True)
            rr += 1
        put(rr, col, "Grand Total", **GT)
        if with_values:
            put(rr, col + 1, ui.fmt(round(sum(data[k] for k in order), 2), "General"), num=True, **GT)
        elif True:
            cells[(rr, col)]["border"] = {"top": ui.ACCENT1}
    reg = {k: REG[k] for k in (["West"] if west else REG_ORDER)}
    pro = WEST_PRO if west else PRO
    if stage == "rows":
        block(1, "Region", REG, REG_ORDER, with_values=False, filt=False)
    elif stage == "p1":
        block(1, "Region", REG, REG_ORDER, filt=False)
    elif stage == "p1f":
        block(1, "Region", REG, REG_ORDER)
    elif stage == "p2":
        block(1, "Region", reg, list(reg))
        block(4, "Product", pro, PRO_ORDER)
    return ui.DictSource(cells, colw={1: 13, 2: 20, 3: 4, 4: 13, 5: 20})


def pivot_overlay_placeholder(v):
    x, y, w, h = v.rect("A3:C18")
    return (f'<div style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;background:#F2F2F2;'
            'border:1px solid #d0d0d0;z-index:15;display:flex;flex-direction:column;align-items:center;'
            'justify-content:center;gap:14px;color:#555;font-size:17px;text-align:center;padding:20px">'
            '<div style="font-size:42px;color:#4F81BD">&#9638;</div><b style="color:#333">PivotTable1</b>'
            '<div>To build a report, choose fields<br>from the PivotTable Field List</div></div>')


PANE_W = 340
def pv(stage, sel="A3", hls=(), pane=True, pane_kw=None, tab="PivotTable Analyze", ctx="PivotTable Analyze",
       rib="PivotTable Analyze", rib_hl=None, rib_kw=None, objs="", dlg="", nb=None, fx=None, west=False):
    src = pivot_source(stage if stage != "blank" else "none", west)
    width = ui.W - PANE_W if pane else ui.W
    v = ui.SheetView(src, zoom=Z_PIV, width=width)
    rr, cc = ui.parse_ref(sel)
    fx = src.formula(rr, cc) if fx is None else fx
    overlay = ""
    if stage == "blank":
        overlay += pivot_overlay_placeholder(v)
    if pane:
        overlay += ui.pivot_pane(**(pane_kw or {"checked": ()}))
    overlay += objs + dlg
    return Frame(v, dict(sel=sel, hls=tuple(hls)),
                 _common("Sheet1", SHEETS1, tab, ctx, rib, rib_hl, rib_kw, nb or sel, fx, None, None, "",
                         int(Z_PIV * 100)), overlay)


PANE_P1 = dict(checked=("Region", "Sales Amount"), rows=("Region",), values=("Sum of Sales Amount",))
PANE_P1F = dict(checked=("Date", "Region", "Sales Amount"), filters=("Date",), rows=("Region",),
                values=("Sum of Sales Amount",))
PANE_P2 = dict(checked=("Date", "Product", "Sales Amount"), filters=("Date",), rows=("Product",),
               values=("Sum of Sales Amount",))


def charts(sel=None, labels=True, title=True, west=False, pie=True):
    cats = ["West"] if west else REG_ORDER
    vals = [REG[c] for c in cats]
    col = ui.column_chart(720, 246, 1170, 372, cats, [int(v) for v in vals],
                          "Region-wise Sales Performance" if title else "Total",
                          selected=(sel == "col"), labels=labels)
    if not pie:
        return col
    pro = WEST_PRO if west else PRO
    p = ui.pie_chart(720, 634, 600, 284, PRO_ORDER, [pro[k] for k in PRO_ORDER],
                     "Product-wise Sales Analysis", selected=(sel == "pie"))
    return col + p


def slicers(west=False, sel_obj=None):
    reg = ui.slicer(1340, 634, 250, "Region", [(k, True) for k in REG_ORDER],
                    selected={"West"} if west else None, did="slr", sel_obj=(sel_obj == "slr"))
    pro = ui.slicer(1610, 634, 250, "Product", [(k, True) for k in PRO_ORDER], did="slp")
    return reg + pro


# ---------------------------------------------------------------- dialogs
def dlg_create_table():
    body = ('<div style="margin:6px 0 2px">Where is the data for your table?</div>'
            f'<div class="dfield">=$A$2:$I${LAST} <span style="float:right;color:#107C41">&#8679;</span></div>'
            '<div><span class="cb on">&#10003;</span>My table has headers</div>')
    return ui.dialog("Create Table", body, 770, 360, 380)


def dlg_create_pivot():
    body = ('<div style="margin:4px 0">Select a table or range</div>'
            '<div style="display:flex;align-items:center;gap:10px">Table/Range: <div class="dfield" style="flex:1">SalesData</div></div>'
            '<div style="margin:8px 0 4px">Choose where you want the PivotTable to be placed</div>'
            '<div style="margin:6px 0"><span class="radio on"></span>New Worksheet</div>'
            '<div style="margin:6px 0"><span class="radio"></span>Existing Worksheet</div>'
            '<div style="display:flex;align-items:center;gap:10px;color:#999">Location: <div class="dfield" style="flex:1;border-color:#ccc">&nbsp;</div></div>')
    return ui.dialog("PivotTable from table or range", body, 720, 300, 480)


def dlg_value_field():
    items = "".join(f'<div class="{"on" if i == "Sum" else ""}">{i}</div>'
                    for i in ["Sum", "Count", "Average", "Max", "Min", "Product"])
    body = ('<div style="margin:4px 0">Source Name: Sales Amount</div>'
            '<div style="display:flex;align-items:center;gap:10px">Custom Name: <div class="dfield" style="flex:1">Sum of Sales Amount</div></div>'
            '<div style="display:flex;gap:18px;border-bottom:1px solid #ddd;margin:6px 0"><b style="border-bottom:2px solid #107C41;padding:4px 0">Summarize Values By</b><span style="padding:4px 0;color:#555">Show Values As</span></div>'
            '<div style="margin:6px 0 2px">Summarize value field by</div>'
            '<div style="color:#555;font-size:12.5px">Choose the type of calculation that you want to use to summarize data from selected field</div>'
            f'<div class="dlist">{items}</div>'
            '<div class="btn" style="display:inline-block;margin-top:6px">Number Format</div>')
    return ui.dialog("Value Field Settings", body, 720, 250, 480)


def dlg_insert_chart():
    left = "".join(f'<div class="{"on" if n == "Column" else ""}" style="padding:7px 12px">{n}</div>'
                   for n in ["Column", "Line", "Pie", "Bar", "Area", "X Y (Scatter)", "Map", "Stock",
                             "Surface", "Radar", "Treemap", "Combo"])
    icons = "".join(f'<div style="width:46px;height:38px;border:{"2px solid #107C41" if i == 0 else "1px solid #ddd"};'
                    f'display:flex;align-items:flex-end;justify-content:center;gap:3px;padding:5px">'
                    f'<i style="width:8px;height:{h1}px;background:#4F81BD"></i><i style="width:8px;height:{h2}px;background:#C0504D"></i></div>'
                    for i, (h1, h2) in enumerate([(18, 24), (14, 26), (26, 26), (20, 12), (16, 24), (12, 22), (24, 18)]))
    prev = "".join(f'<i style="display:inline-block;width:44px;margin:0 12px;height:{h}px;background:#4F81BD;vertical-align:bottom"></i>'
                   for h in (140, 154, 142, 135))
    body = (f'<div style="display:flex;gap:16px;height:390px"><div style="width:160px;border-right:1px solid #e3e3e3">{left}</div>'
            f'<div style="flex:1"><div style="display:flex;gap:8px;margin:4px 0 10px">{icons}</div>'
            '<b>Clustered Column</b><div style="border:1px solid #e3e3e3;margin-top:10px;height:280px;display:flex;align-items:flex-end;justify-content:center;padding:20px">'
            f'{prev}</div></div></div>')
    return ui.dialog("Insert Chart", body, 520, 220, 880)


def dlg_insert_slicers():
    fields = ["Date", "Invoice No", "Region", "Salesperson", "Product", "Qty", "Unit Price",
              "Sales Amount", "Payment Status"]
    body = "".join(f'<div style="margin:6px 0"><span class="cb{" on" if f in ("Region", "Product") else ""}">'
                   f'{"&#10003;" if f in ("Region", "Product") else ""}</span>{f}</div>' for f in fields)
    return ui.dialog("Insert Slicers", body, 820, 270, 300)


def dlg_report_connections():
    body = ('<div style="margin:4px 0 8px">Select PivotTable and PivotChart reports to connect to this filter</div>'
            '<div style="border:1px solid #bdbdbd"><div style="display:flex;background:#f3f3f3;padding:5px 10px;font-weight:600">'
            '<span style="width:40px"></span><span style="flex:1">Name</span><span style="width:110px">Sheet</span></div>'
            '<div style="display:flex;padding:6px 10px"><span style="width:40px"><span class="cb on">&#10003;</span></span>'
            '<span style="flex:1">PivotTable1</span><span style="width:110px">Sheet1</span></div>'
            '<div style="display:flex;padding:6px 10px"><span style="width:40px"><span class="cb on">&#10003;</span></span>'
            '<span style="flex:1">PivotTable2</span><span style="width:110px">Sheet1</span></div>'
            '<div style="height:90px"></div></div>')
    return ui.dialog("Report Connections (Region)", body, 720, 300, 480)


# ================================================================== narration
# Each beat: caption (romanized Hinglish, shown on screen), tts (Devanagari
# Hinglish for the voice), frames [(frac, Frame)], cursor [(frac, target, click)],
# sfx [(frac, kind)], typing (target, text, f0, f1) optional.
# frac is a fraction of the spoken narration (0 = speech start, 1 = speech end).
BEATS = []
SCENE = [0, ""]


def scene(n, badge):
    SCENE[0], SCENE[1] = n, badge


def B(caption, tts, frames, cursor=(), sfx=(), typing=None, extra=0.0):
    if isinstance(frames, Frame):
        frames = [(0.0, frames)]
    BEATS.append(dict(scene=SCENE[0], badge=SCENE[1], caption=caption, tts=tts,
                      frames=list(frames), cursor=list(cursor), sfx=list(sfx), typing=typing, extra=extra))


def typed(make, text, f0, f1, steps=None):
    """Frames that type `text` progressively between f0 and f1. Returns (frames, key_fracs)."""
    n = len(text)
    steps = steps or min(14, max(4, n // 6))
    out, keys = [], []
    for i in range(1, steps + 1):
        k = round(n * i / steps)
        f = f0 + (f1 - f0) * (i - 1) / steps
        out.append((f, make(text[:k])))
    per = (f1 - f0) / max(n, 1)
    keys = [f0 + per * j for j in range(n) if text[j] != " " or j % 2 == 0][::2]
    return out, keys


T_SA = "=[@Qty]*[@[Unit Price]]"
F_B6 = DSH.formula(6, 2)

# -------------------------------- Scene 1 : Introduction
scene(1, "Scene 1 &bull; Introduction")
D0 = dash("A1")
B("Namaskaar doston, aur swagat hai aapka is channel par.",
  "नमस्कार दोस्तों, और स्वागत है आपका इस चैनल पर।", D0,
  cursor=[(0.6, "cell:B6", False)])
B("Aaj hum banayenge ek complete, professional <b>Sales MIS Report</b> — bilkul usi tarah jaise companies mein MIS executives banate hain.",
  "आज हम बनाएंगे एक complete, professional Sales MIS Report — बिल्कुल उसी तरह, जैसे companies में MIS executives बनाते हैं।",
  D0, cursor=[(0.5, "cell:B2", False)])
B("<b>MIS</b> ka matlab hai <b>Management Information System</b> — raw data ko aise summarise karna ki management ek nazar mein decision le sake.",
  "MIS का मतलब है management information system — यानी raw data को ऐसे summarise करना, कि management एक नज़र में decision ले सके।",
  D0, cursor=[(0.3, "cell:D6", False), (0.8, "cell:H6", False)])
B("Video ke end tak aapke paas hoga ek live dashboard — KPIs, region aur product analysis, Pivot Tables, charts aur slicers.",
  "इस video के end तक आपके पास होगा एक live dashboard — KPIs, region और product analysis, pivot tables, charts और slicers।",
  dash("A1", hls=("B11:D17", "F11:H16")), cursor=[(0.35, "cell:C14", False), (0.75, "cell:G14", False)])
B("Main sirf dikhaunga nahi, har formula samjhaunga bhi. Toh chaliye shuru karte hain!",
  "मैं सिर्फ़ दिखाऊंगा नहीं, हर formula समझाऊंगा भी। तो चलिए, शुरू करते हैं!", D0,
  cursor=[(0.7, "id:sheet-Raw_Data", False)])

# -------------------------------- Scene 2 : Raw data
scene(2, "Scene 2 &bull; Raw Data")
RPRE = dict(table=False, blank=H_BLANK)
B("Sabse pehle data samajhte hain. Yeh <b>Raw_Data</b> sheet hai — isme October 2026 ki <b>180 transactions</b> hain.",
  "सबसे पहले data को समझते हैं। यह Raw Data sheet है, और इसमें October दो हज़ार छब्बीस की, कुल एक सौ अस्सी transactions हैं।",
  [(0.0, D0), (0.12, raw("A3", hls=("A2:I2",), **RPRE))],
  cursor=[(0.08, "id:sheet-Raw_Data", True), (0.5, "cell:A3", True)])
B("<b>Date</b> invoice ki tareekh hai, <b>Invoice No</b> har bill ka unique number, aur <b>Region</b> — East, West, North ya South.",
  "Date है invoice की तारीख़, Invoice number है हर bill का unique number, और Region — East, West, North या South।",
  raw("C3", hls=("A2:C29",), **RPRE), cursor=[(0.1, "cell:A3", False), (0.4, "cell:B3", False), (0.8, "cell:C3", True)])
B("Phir <b>Salesperson</b>, <b>Product</b> — Tea A, Tea B ya Tea C, <b>Qty</b> yaani kitne units bike, aur <b>Unit Price</b>.",
  "फिर Salesperson, Product — Tea A, Tea B या Tea C, Quantity यानी कितने units बिके, और Unit Price।",
  raw("F3", hls=("D2:G29",), **RPRE), cursor=[(0.2, "cell:D3", False), (0.5, "cell:E3", False), (0.85, "cell:F3", True)])
B("<b>Sales Amount</b> abhi khaali hai — ise hum formula se bharenge. Aur <b>Payment Status</b> batata hai — Paid ya Pending.",
  "Sales Amount का column अभी ख़ाली है — इसे हम formula से भरेंगे। और Payment Status बताता है — Paid या Pending।",
  raw("I3", hls=("H2:I29",), **RPRE), cursor=[(0.2, "cell:H3", False), (0.75, "cell:I3", True)])
B("Shortcut: <b>Ctrl + ↓</b> — cursor seedha aakhri row <b>182</b> par. Yaani header ke neeche <b>180 data rows</b>.",
  "एक useful shortcut — control और down arrow दबाइए। Cursor सीधा आख़िरी row, यानी एक सौ बयासी पर पहुंच जाता है। मतलब header के नीचे कुल एक सौ अस्सी data rows।",
  [(0.0, raw("A3", **RPRE)), (0.3, raw(f"A{LAST}", top=LAST - 22, hls=(f"A{LAST}:I{LAST}",), **RPRE))],
  cursor=[(0.1, "cell:A3", True)], sfx=[(0.25, "key"), (0.28, "key")])
B("Important: <b>ek row = ek unique invoice</b>. Har Invoice No bhara hua hai aur kabhi repeat nahi hota.",
  "एक बहुत important बात — इस data में एक row का मतलब है एक unique invoice। हर Invoice number भरा हुआ है, और कभी repeat नहीं होता।",
  [(0.0, raw("B3", hls=("B3:B29",), **RPRE))], cursor=[(0.05, "cell:B3", True), (0.7, "cell:B20", False)],
  sfx=[(0.02, "key"), (0.04, "key")])
B("Yeh assumption humne <b>Read_Me</b> sheet mein likh rakhi hai — aage Total Orders isi par based hoga.",
  "यह assumption हमने Read Me sheet में लिख रखी है, और आगे Total Orders का calculation इसी पर based होगा।",
  [(0.0, raw("B3", **RPRE)), (0.15, readme("B16", hls=("B14:B19",)))],
  cursor=[(0.1, "id:sheet-Read_Me", True), (0.6, "cell:B16", False)])

# -------------------------------- Scene 3 : Excel Table
scene(3, "Scene 3 &bull; Excel Table")
CT = dlg_create_table()
B("Ab is data ko <b>Excel Table</b> mein convert karte hain — formulas readable bante hain aur range apne aap expand hoti hai.",
  "अब इस data को एक Excel Table में convert करते हैं। इससे formulas readable बनते हैं, और नया data आने पर range अपने आप expand हो जाती है।",
  [(0.0, readme("B16")), (0.1, raw("A3", **RPRE))],
  cursor=[(0.06, "id:sheet-Raw_Data", True), (0.6, "cell:C5", True)])
B("Data ke kisi bhi cell par click karke <b>Ctrl + T</b> dabaiye.",
  "Data के किसी भी cell पर click करके, control plus T दबाइए।",
  [(0.0, raw("C5", **RPRE)), (0.75, raw("C5", rng=f"A2:I{LAST}", ants=f"A2:I{LAST}", dlg=CT, **RPRE))],
  cursor=[(0.2, "cell:C5", True)], sfx=[(0.68, "key"), (0.71, "key"), (0.76, "pop")])
B(f"<b>Create Table</b> mein range <b>=$A$2:$I${LAST}</b> aa gayi. 'My table has headers' tick hona chahiye — phir <b>OK</b>.",
  "Create Table box में range अपने आप आ गई — ए टू से आई वन एटी टू तक। ध्यान रखिए, My table has headers tick होना चाहिए। फिर OK दबाइए।",
  [(0.0, raw("C5", rng=f"A2:I{LAST}", ants=f"A2:I{LAST}", dlg=CT, **RPRE)),
   (0.94, raw("C5", tab="Table Design", ctx="Table Design", blank=H_BLANK))],
  cursor=[(0.55, "id:dlg-ok", False), (0.92, "id:dlg-ok", True)])
B("Table ban gaya — <b>banded rows</b>, har header par <b>filter button</b>, aur ribbon mein <b>Table Design</b> tab.",
  "Table बन गया! देखिए — banded rows आ गईं, हर header पर filter button आ गया, और ribbon में Table Design tab भी खुल गया।",
  raw("C5", tab="Table Design", ctx="Table Design", blank=H_BLANK, hls=("A2:I2",)),
  cursor=[(0.35, "cell:C2", False), (0.8, "id:tab-Table Design", False)])
tn_frames, tn_keys = typed(lambda s: raw("C5", tab="Table Design", ctx="Table Design", blank=H_BLANK,
                                         rib_hl="rb-tablename", rib_kw={"table_name": s or "Table1"}),
                           "SalesData", 0.25, 0.55, steps=5)
B("<b>Table Name</b> box mein likhiye <b>SalesData</b> aur Enter. Yahi naam aage har formula mein use hoga.",
  "अब Table Name box में नाम लिखिए — Sales Data — और enter दबाइए। यही नाम आगे हर formula में use होगा।",
  [(0.0, raw("C5", tab="Table Design", ctx="Table Design", blank=H_BLANK, rib_hl="rb-tablename"))] + tn_frames,
  cursor=[(0.18, "id:rb-tablename", True)], sfx=[(f, "key") for f in tn_keys] + [(0.58, "key")])

# -------------------------------- Scene 4 : Sales Amount
scene(4, "Scene 4 &bull; Sales Amount")
B("Ab Sales Amount. Table ke andar formulas <b>structured references</b> use karte hain — cell address nahi, column ka naam.",
  "अब Sales Amount निकालते हैं। Table के अंदर formulas structured references use करते हैं — यानी cell address नहीं, column का नाम।",
  raw("H3", blank=H_BLANK, hls=("H2:H29",)), cursor=[(0.3, "cell:H3", True)])
sa_frames, sa_keys = typed(lambda s: raw("H3", blank=H_BLANK, edit=s), T_SA, 0.12, 0.85, steps=10)
B("<b>H3</b> mein type kijiye: <b>=[@Qty]*[@[Unit Price]]</b>",
  "H three cell में type कीजिए — equals, bracket, at, Qty, close bracket, star, फिर bracket, at, bracket, Unit Price, और brackets close।",
  [(0.0, raw("H3", blank=H_BLANK))] + sa_frames, cursor=[(0.05, "cell:H3", False)],
  sfx=[(f, "key") for f in sa_keys])
B("<b>@</b> ka matlab — <b>isi row</b> ki value. Unit Price do shabdon ka naam hai, isliye uske around <b>double brackets</b>.",
  "यहां at का मतलब है — इसी row की value। और Unit Price दो शब्दों का नाम है, इसलिए उसके around double brackets लगते हैं।",
  raw("H3", blank=H_BLANK, edit=T_SA, hls=("F2:G2",)), cursor=[(0.3, "cell:F2", False), (0.7, "cell:G2", False)])
B("<b>Enter</b> dabaate hi formula <b>poore column</b> mein apne aap fill — saari 180 rows ek saath.",
  "Enter दबाते ही, Excel यह formula पूरे column में अपने आप भर देता है — सारी एक सौ अस्सी rows एक साथ।",
  [(0.0, raw("H3", blank=H_BLANK, edit=T_SA)), (0.08, raw("H4", hls=("H3:H29",)))],
  cursor=[(0.5, "cell:H12", False)], sfx=[(0.05, "key")])
B("Check: row 3 mein <b>55 × 120.00 = 6,600.00</b> ✔. Column select karein toh status bar mein <b>Sum: 1,428,190.00</b>.",
  "Check कर लीजिए — row तीन में पचपन गुणा एक सौ बीस, बराबर छह हज़ार छह सौ। और पूरा column select करें, तो status bar में Sum दिखता है — चौदह लाख अट्ठाईस हज़ार एक सौ नब्बे।",
  [(0.0, raw("H3", hls=("F3:H3",))),
   (0.5, raw("H3", rng=f"H3:H{LAST}", info="Average: 7,934.39&nbsp;&nbsp;&nbsp;Count: 180&nbsp;&nbsp;&nbsp;Sum: 1,428,190.00"))],
  cursor=[(0.08, "cell:H3", True), (0.85, (1500, 1064), False)],
  sfx=[(0.46, "key"), (0.48, "key"), (0.5, "key")])

# -------------------------------- Scene 5 : KPIs
scene(5, "Scene 5 &bull; KPI Summary")
B("Ab dashboard ka dil — <b>KPI cards</b>. Yaad rakhiye, final values kabhi haath se type nahi karte.",
  "अब आता है dashboard का दिल — KPI cards। याद रखिए, final values हम कभी हाथ से type नहीं करते।",
  [(0.0, raw("H4")), (0.12, dash("A1"))], cursor=[(0.08, "id:sheet-MIS_Dashboard", True), (0.7, "cell:D6", False)])
B("Sab kuch formula se aata hai — data badla, toh number apne aap update.",
  "सब कुछ formula से आता है, ताकि data बदलते ही number अपने आप update हो जाए।",
  dash("A1", hls=("B5:I6", "B8:C9")), cursor=[(0.5, "cell:F6", False)])
B6_BLANK = {(6, 2)}
b6_frames, b6_keys = typed(lambda s: dash("B6", edit=s, blank=B6_BLANK), F_B6, 0.18, 0.92, steps=14)
B("<b>Total Sales (B6)</b>: <b>SUMIFS</b> sirf October ka Sales Amount jodta hai — date range <b>DATE()</b> se di hai.",
  "पहला KPI — Total Sales। B six में sum ifs formula है, जो सिर्फ़ October का Sales Amount जोड़ता है। Date की range हमने DATE function से दी है — एक October से बड़ा या बराबर, और एक November से छोटा।",
  [(0.0, dash("B6", blank=B6_BLANK))] + b6_frames, cursor=[(0.1, "cell:B6", True)],
  sfx=[(f, "key") for f in b6_keys])
B("Result: <b>1,428,190.00</b> — yaani chaudah lakh, atthais hazaar, ek sau nabbe.",
  "Result आता है — चौदह लाख, अट्ठाईस हज़ार, एक सौ नब्बे।",
  [(0.0, dash("B6", edit=F_B6, blank=B6_BLANK)), (0.05, dash("B6", hls=("B6",)))],
  cursor=[(0.15, "cell:B6", False)],          # after Enter, move off the value so it stays readable
  sfx=[(0.03, "key"), (0.4, "chime")])
B("<b>Total Orders (D6)</b>: ek row = ek invoice, isliye <b>COUNTIFS</b> se October ke non-blank invoices — <b>180</b>.",
  "दूसरा — Total Orders। क्योंकि एक row एक unique invoice है, हम count ifs से October के non blank invoices गिनते हैं। Total — एक सौ अस्सी।",
  dash("D6", hls=("D6",)), cursor=[(0.1, "cell:D6", True)])
B("<b>Total Units Sold (F6)</b>: wahi date condition, Qty ka SUMIFS — <b>11,834</b>.",
  "तीसरा — Total Units Sold। वही date condition, लेकिन Quantity का sum ifs। जवाब — ग्यारह हज़ार आठ सौ चौंतीस units।",
  dash("F6", hls=("F6",)), cursor=[(0.1, "cell:F6", True)])
B("<b>Average Sales / Order (H6)</b>: <b>=IFERROR(B6/D6,0)</b> — orders zero hon toh bhi error nahi. Result <b>7,934.39</b>.",
  "चौथा — Average Sales per Order। Total Sales को Total Orders से divide किया, और if error में लपेट दिया, ताकि कभी orders zero हों तो error न आए। जवाब — लगभग सेवन थाउज़ंड नाइन हंड्रेड थर्टी फ़ोर।",
  dash("H6", hls=("H6",)), cursor=[(0.1, "cell:H6", True)])
B("<b>Pending Sales Amount (B9)</b>: Payment Status = \"Pending\" wala SUMIFS — <b>372,985.00</b>.",
  "पांचवां — Pending Sales Amount। इसमें sum ifs की एक extra condition है — Payment Status, Pending। Pending Sales Amount है — थ्री लाख, सेवेंटी टू थाउज़ंड, नाइन हंड्रेड एटी फ़ाइव।",
  dash("B9", hls=("B9",)), cursor=[(0.1, "cell:B9", True)])
B("Naam jaan-boojhkar <b>Pending Sales Amount</b> rakha hai, 'Outstanding Receivables' nahi — data mein payment received ya balance ki field hi nahi hai.",
  "ध्यान दीजिए — इसका नाम हमने जान-बूझकर Pending Sales Amount रखा है, Outstanding Receivables नहीं। क्योंकि इस data में payment received या balance की कोई field ही नहीं है।",
  dash("B9", hls=("B8",)), cursor=[(0.3, "cell:B8", False)])

# -------------------------------- Scene 6 : Region / Product
scene(6, "Scene 6 &bull; Region Analysis")
B("Ab <b>region-wise</b> analysis — management dekhna chahta hai kaunsa region aage hai.",
  "अब region wise analysis। Management देखना चाहता है कि कौन सा region सबसे आगे है।",
  dash("B13", hls=("B11:D17",)), cursor=[(0.3, "cell:B13", True)])
B("<b>C13</b> mein wahi SUMIFS + ek condition: <b>SalesData[Region] = B13</b>. Region cell se aata hai, isliye formula neeche copy ho jaata hai.",
  "C thirteen में वही sum ifs है, बस एक extra condition के साथ — Region बराबर B thirteen। Region का नाम cell से आता है, इसलिए यही formula नीचे copy हो जाता है।",
  [(0.0, dash("C13")), (0.72, dash("C13", rng="C13:C16"))],
  cursor=[(0.08, "cell:C13", True), (0.7, "cell:C16", False)], sfx=[(0.74, "key"), (0.77, "key")])
B("East <b>351,030</b> • West <b>383,675</b> (sabse strong) • North <b>355,130</b> • South <b>338,355</b>.",
  "East — थ्री लाख, फ़िफ़्टी वन थाउज़ंड, थर्टी। West — थ्री लाख, एटी थ्री थाउज़ंड, सिक्स हंड्रेड सेवेंटी फ़ाइव, यह सबसे strong region है। North — थ्री लाख, फ़िफ़्टी फ़ाइव थाउज़ंड, वन हंड्रेड थर्टी। और South — थ्री लाख, थर्टी एट थाउज़ंड, थ्री हंड्रेड फ़िफ़्टी फ़ाइव।",
  [(0.0, dash("C13", hls=("B13:C13",))), (0.24, dash("C14", hls=("B14:C14",))),
   (0.6, dash("C15", hls=("B15:C15",))), (0.8, dash("C16", hls=("B16:C16",)))],
  cursor=[(0.02, "cell:C13", True), (0.24, "cell:C14", True), (0.6, "cell:C15", True), (0.8, "cell:C16", True)])
B("Total <b>(C17)</b> = <b>1,428,190.00</b> — bilkul Total Sales ke barabar. Yahi hamara <b>cross-check</b> hai ✔",
  "और इन चारों का total — C seventeen में — बिल्कुल Total Sales के बराबर है। यही हमारा cross check है कि numbers reconcile हो रहे हैं।",
  dash("C17", hls=("C17", "B6")), cursor=[(0.1, "cell:C17", True), (0.6, "cell:B6", False)], sfx=[(0.75, "chime")])
B("Isi tarah <b>Product-wise</b>: Tea A 385,320 • Tea B 432,820 • Tea C <b>610,050</b> — total phir se <b>1,428,190.00</b> ✔",
  "इसी तरह product wise — Tea A, Tea B, और Tea C। Tea C सबसे ज़्यादा बिकी, और इनका total भी फिर से वही — चौदह लाख अट्ठाईस हज़ार एक सौ नब्बे।",
  dash("G16", hls=("F11:H16",)), cursor=[(0.2, "cell:G15", False), (0.6, "cell:G16", True)], sfx=[(0.85, "chime")])

# -------------------------------- Scene 7 : Pivot Table
scene(7, "Scene 7 &bull; Pivot Table")
B("Formulas ke saath <b>Pivot Table</b> bhi seekhte hain — drag-and-drop se instant summary.",
  "Formulas के साथ-साथ, अब pivot table भी सीखते हैं — drag and drop से instant summary।",
  [(0.0, dash("G16")), (0.12, raw("C5")), (0.6, raw("C5", tab="Insert"))],
  cursor=[(0.08, "id:sheet-Raw_Data", True), (0.35, "cell:C5", True), (0.58, "id:tab-Insert", True)])
B("SalesData ke kisi cell par click → <b>Insert</b> → <b>PivotTable</b> → From Table/Range.",
  "Sales Data table के किसी cell पर click कीजिए, फिर Insert tab में, PivotTable, From Table or Range।",
  [(0.0, raw("C5", tab="Insert", rib_hl="rb-pivot")), (0.85, raw("C5", tab="Insert", dlg=dlg_create_pivot()))],
  cursor=[(0.55, "id:rb-pivot", False), (0.82, "id:rb-pivot", True)], sfx=[(0.86, "pop")])
B("Table/Range mein <b>SalesData</b> already aa gaya. <b>New Worksheet</b> choose karke <b>OK</b>.",
  "Table or Range में Sales Data पहले से आ गया है। New Worksheet choose कीजिए, और OK दबाइए।",
  [(0.0, raw("C5", tab="Insert", dlg=dlg_create_pivot())), (0.95, pv("blank", pane_kw={"checked": ()}))],
  cursor=[(0.5, "id:dlg-ok", False), (0.93, "id:dlg-ok", True)])
B("Nayi sheet bani — right side <b>PivotTable Fields</b>. <b>Region</b> ko <b>Rows</b> mein, <b>Sales Amount</b> ko <b>Values</b> mein drag kijiye.",
  "नई sheet बन गई, और right side में PivotTable Fields pane खुल गया। अब Region को Rows area में drag कीजिए, और Sales Amount को Values area में।",
  [(0.0, pv("blank", pane_kw={"checked": ()})),
   (0.62, pv("rows", pane_kw=dict(checked=("Region",), rows=("Region",)))),
   (0.95, pv("p1", sel="B4", pane_kw=PANE_P1))],
  cursor=[(0.45, "id:pf-Region", True), (0.6, "id:pa-rows", False), (0.8, "id:pf-Sales Amount", True), (0.94, "id:pa-values", False)])
B("Dhyaan: Values mein <b>Sum of Sales Amount</b> hona chahiye. 'Count of' dikhe toh → <b>Value Field Settings</b> → <b>Sum</b>.",
  "ध्यान दीजिए — Values में Sum of Sales Amount होना चाहिए। अगर कभी Count of दिख जाए, तो Value Field Settings में जाकर Sum select कीजिए।",
  [(0.0, pv("p1", sel="B4", pane_kw=PANE_P1, rib_hl="rb-fieldsettings")),
   (0.42, pv("p1", sel="B4", pane_kw=PANE_P1, dlg=dlg_value_field())),
   (0.97, pv("p1", sel="B4", pane_kw=PANE_P1))],
  cursor=[(0.2, "id:rb-fieldsettings", False), (0.38, "id:rb-fieldsettings", True), (0.95, "id:dlg-ok", True)],
  sfx=[(0.43, "pop")])
B("<b>Date</b> ko <b>Filters</b> mein daaliye, October 2026 rakhiye. Grand Total = <b>1428190</b> — dashboard se match ✔",
  "अब Date को Filters area में डालिए, और October दो हज़ार छब्बीस select कीजिए। Grand Total देखिए — बिल्कुल dashboard के Total Sales से match करता है।",
  [(0.0, pv("p1", sel="B4", pane_kw=PANE_P1)), (0.3, pv("p1f", sel="B1", pane_kw=PANE_P1F)),
   (0.65, pv("p1f", sel="B8", pane_kw=PANE_P1F, hls=("A8:B8",)))],
  cursor=[(0.15, "id:pf-Date", True), (0.28, "id:pa-filters", False), (0.62, "cell:B8", True)], sfx=[(0.8, "chime")])
B("Isi tarah doosra pivot — <b>Product</b> Rows mein, Sales Amount Values mein.",
  "इसी तरह एक दूसरा pivot बनाइए — Product को Rows में, और Sales Amount को Values में।",
  [(0.0, pv("p1f", sel="B8", pane_kw=PANE_P1F)),
   (0.5, pv("p2", sel="D4", pane_kw=PANE_P2, rib_kw={"pt_name": "PivotTable2"}))],
  cursor=[(0.3, "id:pf-Product", True), (0.48, "id:pa-rows", False)])

# -------------------------------- Scene 8 : Charts
scene(8, "Scene 8 &bull; Charts")
P2 = dict(pane=False)
B("Ab charts — visuals se management ko baat turant samajh aati hai.",
  "अब बनाते हैं charts — क्योंकि visuals से management को बात तुरंत समझ आती है।",
  [(0.0, pv("p2", sel="D4", pane_kw=PANE_P2)), (0.3, pv("p2", sel="A4", rib_hl="rb-fieldlist", **P2))],
  cursor=[(0.25, "id:rb-fieldlist", True), (0.8, "cell:A4", True)])
B("Region pivot select → <b>PivotTable Analyze</b> → <b>PivotChart</b> → <b>Clustered Column</b> → OK.",
  "Region वाले pivot पर click कीजिए, PivotTable Analyze में PivotChart, और Clustered Column choose करके OK।",
  [(0.0, pv("p2", sel="A4", rib_hl="rb-pivotchart2", **P2)), (0.4, pv("p2", sel="A4", dlg=dlg_insert_chart(), **P2)),
   (0.97, pv("p2", sel="A4", objs=charts(sel="col", labels=False, title=False, pie=False), tab="Design",
             ctx="Chart Design", rib="Chart Design", nb="Chart 1", fx="", **P2))],
  cursor=[(0.3, "id:rb-pivotchart2", True), (0.94, "id:dlg-ok", True)], sfx=[(0.41, "pop")])
B("Chart aa gaya. Right-click → <b>Add Data Labels</b> — har bar par exact value. Title: <b>Region-wise Sales Performance</b>.",
  "Chart आ गया। अब bars पर right click करके Add Data Labels — ताकि हर bar पर exact value दिखे। और title रखिए — Region wise Sales Performance।",
  [(0.0, pv("p2", sel="A4", objs=charts(sel="col", labels=False, title=False, pie=False), tab="Design",
            ctx="Chart Design", rib="Chart Design", nb="Chart 1", fx="", rib_hl="rb-addelem", **P2)),
   (0.42, pv("p2", sel="A4", objs=charts(sel="col", labels=True, title=False, pie=False), tab="Design",
             ctx="Chart Design", rib="Chart Design", nb="Chart 1", fx="", **P2)),
   (0.75, pv("p2", sel="A4", objs=charts(sel="col", labels=True, title=True, pie=False), tab="Design",
             ctx="Chart Design", rib="Chart Design", nb="Chart 1", fx="", **P2))],
  cursor=[(0.25, (943, 470), False), (0.4, (943, 470), True), (0.72, (1305, 287), True)])
B("Product pivot ke liye <b>Pie chart</b> — <b>Product-wise Sales Analysis</b>. Tea C ka share sabse bada.",
  "Product वाले pivot के लिए, हम Pie chart बनाएंगे। Title रखिए — Product wise Sales Analysis। देखिए, Tea C का share सबसे बड़ा है।",
  [(0.0, pv("p2", sel="D4", objs=charts(), **P2)),
   (0.2, pv("p2", sel="D4", objs=charts(sel="pie"), tab="Design", ctx="Chart Design", rib="Chart Design",
            nb="Chart 2", fx="", **P2))],
  cursor=[(0.15, "cell:D4", True), (0.6, (948, 793), True)])

# -------------------------------- Scene 9 : Slicers
scene(9, "Scene 9 &bull; Slicers")
CH = charts()
B("Ab report ko <b>interactive</b> banate hain — <b>slicers</b> se.",
  "अब इस report को interactive बनाते हैं — slicers की मदद से।",
  pv("p2", sel="A4", objs=CH, **P2), cursor=[(0.5, "cell:A5", True)])
B("Pivot select → <b>PivotTable Analyze</b> → <b>Insert Slicer</b> → <b>Region</b> aur <b>Product</b> tick → OK.",
  "Pivot पर click कीजिए, PivotTable Analyze में Insert Slicer, फिर Region और Product tick करके OK।",
  [(0.0, pv("p2", sel="A5", objs=CH, rib_hl="rb-insslicer", **P2)),
   (0.35, pv("p2", sel="A5", objs=CH, dlg=dlg_insert_slicers(), **P2)),
   (0.95, pv("p2", sel="A5", objs=CH + slicers(sel_obj="slr"), tab="Slicer", ctx="Slicer", rib="Slicer", nb="Region", fx="", **P2))],
  cursor=[(0.3, "id:rb-insslicer", True), (0.93, "id:dlg-ok", True)], sfx=[(0.36, "pop")])
SL = pv("p2", sel="A5", objs=CH + slicers(sel_obj="slr"), tab="Slicer", ctx="Slicer", rib="Slicer", nb="Region", fx="", **P2)
B("Ek slicer se dono pivots: slicer par right-click → <b>Report Connections</b> → dono PivotTables tick.",
  "एक ही slicer से दोनों pivots control करने के लिए, slicer पर right click कीजिए, Report Connections में जाइए, और दोनों PivotTables tick कर दीजिए।",
  [(0.0, pv("p2", sel="A5", objs=CH + slicers(sel_obj="slr"), tab="Slicer", ctx="Slicer", rib="Slicer",
            rib_hl="rb-reportconn", nb="Region", fx="", **P2)),
   (0.45, pv("p2", sel="A5", objs=CH + slicers(sel_obj="slr"), tab="Slicer", ctx="Slicer", rib="Slicer",
             nb="Region", fx="", dlg=dlg_report_connections(), **P2)),
   (0.97, SL)],
  cursor=[(0.38, "id:rb-reportconn", True), (0.94, "id:dlg-ok", True)], sfx=[(0.46, "pop")])
B("Yaad rakhiye: slicer sirf <b>same data source</b> (SalesData) wale pivots ko jod sakta hai.",
  "याद रखिए — slicer सिर्फ़ उन्हीं pivots को जोड़ सकता है, जो same data source, यानी Sales Data से बने हैं।",
  SL, cursor=[(0.5, "id:slr", False)])
WEST = pv("p2", sel="A5", objs=charts(west=True) + slicers(west=True), west=True, **P2)
B("Ab <b>West</b> par click — dono pivots aur charts turant West ke liye filter ho gaye!",
  "अब West पर click कीजिए — देखिए, दोनों pivots और charts तुरंत West के लिए filter हो गए।",
  [(0.0, SL), (0.3, WEST)], cursor=[(0.27, "id:slr-West", True), (0.8, "cell:B5", False)])
B("Filter hatane ke liye slicer ka <b>Clear Filter</b> button — totals wapas poore.",
  "Filter हटाने के लिए slicer का Clear Filter button दबाइए — totals वापस पूरे हो जाते हैं।",
  [(0.0, WEST), (0.4, pv("p2", sel="A5", objs=CH + slicers(), **P2))],
  cursor=[(0.37, "id:slr-clear", True)])

# -------------------------------- Scene 10 : Formatting
scene(10, "Scene 10 &bull; Formatting")
B("Ab finishing — professional dashboard ka matlab decoration nahi, <b>clarity</b>.",
  "अब finishing touch — formatting। Professional dashboard का मतलब ज़्यादा decoration नहीं, बल्कि clarity है।",
  [(0.0, pv("p2", sel="A5", objs=CH + slicers(), **P2)), (0.15, dash("A1", grid=True, sheets=SHEETS1))],
  cursor=[(0.1, "id:sheet-MIS_Dashboard", True), (0.7, "cell:E2", False)])
B("Bold title, KPI cards ka consistent colour, aur numbers par <b>thousands separator</b> — <b>Ctrl + Shift + 1</b>.",
  "Title bold और बड़ा, KPI cards का एक consistent colour, और numbers पर thousands separator — shortcut है, control shift one।",
  dash("B6", rng="B6:I6", grid=True, sheets=SHEETS1, rib_hl="rb-comma", rib_kw={"numfmt": "Number"}),
  cursor=[(0.3, "cell:B6", True), (0.6, "id:rb-comma", False)], sfx=[(0.8, "key"), (0.82, "key"), (0.84, "key")])
B("<b>View</b> tab → <b>Gridlines</b> untick — look ekdum clean.",
  "View tab में जाकर Gridlines का tick हटा दीजिए — look एकदम clean हो जाता है।",
  [(0.0, dash("A1", grid=True, sheets=SHEETS1)),
   (0.2, dash("A1", grid=True, sheets=SHEETS1, tab="View", rib="View", rib_kw={"gridlines": True}, rib_hl="rb-gridlines")),
   (0.62, dash("A1", grid=False, sheets=SHEETS1, tab="View", rib="View", rib_kw={"gridlines": False}, rib_hl="rb-gridlines"))],
  cursor=[(0.15, "id:tab-View", True), (0.6, "id:rb-gridlines", True)])
B("Ek font, consistent spacing, do-teen colours — report manager ke liye hai, <b>readability</b> sabse upar.",
  "एक font, consistent spacing, और बस दो-तीन colours। Report manager के लिए है, इसलिए readability सबसे ऊपर।",
  dash("A1", sheets=SHEETS1, tab="View", rib="View", rib_kw={"gridlines": False}),
  cursor=[(0.4, "cell:E3", False), (0.85, "cell:E9", False)])

# -------------------------------- Scene 11 : Refresh & validate
scene(11, "Scene 11 &bull; Refresh &amp; Validate")
B("Report tab tak bharosemand nahi jab tak <b>refresh</b> aur <b>validation</b> na ho.",
  "कोई भी MIS report तब तक भरोसेमंद नहीं, जब तक उसका refresh और validation न हो।",
  dash("A1", sheets=SHEETS1), cursor=[(0.6, "cell:F3", False)])
B("Raw_Data mein nayi row aaye — Table hone se SUMIFS / COUNTIFS use <b>apne aap</b> pakad lete hain.",
  "अगर Raw Data में नई row add होती है, तो table होने की वजह से sum ifs और count ifs उसे अपने आप पकड़ लेते हैं।",
  [(0.0, dash("A1", sheets=SHEETS1)), (0.15, raw(f"A{LAST+1}", top=LAST - 20, sheets=SHEETS1, hls=(f"A{LAST}:I{LAST+1}",)))],
  cursor=[(0.1, "id:sheet-Raw_Data", True), (0.6, f"cell:A{LAST+1}", True)])
B("Pivot Tables ko refresh karna padta hai: <b>Data</b> → <b>Refresh All</b> — ya <b>Ctrl + Alt + F5</b>.",
  "लेकिन pivot tables को refresh करना पड़ता है — Data tab में Refresh All, या shortcut — control alt F five।",
  [(0.0, raw(f"A{LAST+1}", top=LAST - 20, sheets=SHEETS1)),
   (0.2, raw(f"A{LAST+1}", top=LAST - 20, sheets=SHEETS1, tab="Data", rib="Data", rib_hl="rb-refreshall"))],
  cursor=[(0.15, "id:tab-Data", True), (0.6, "id:rb-refreshall", True)], sfx=[(0.88, "key"), (0.9, "key"), (0.92, "key")])
B("Teen cross-check: Region total = Product total = Pivot Grand Total = <b>Total Sales 1,428,190</b>.",
  "और validation के लिए तीन cross check — region total, product total, और pivot का Grand Total, तीनों Total Sales के बराबर होने चाहिए।",
  [(0.0, raw(f"A{LAST+1}", top=LAST - 20, sheets=SHEETS1)), (0.1, dash("C17", sheets=SHEETS1, hls=("C17", "G16", "B6")))],
  cursor=[(0.06, "id:sheet-MIS_Dashboard", True), (0.4, "cell:C17", False), (0.7, "cell:G16", False)])
B("Teeno match ✔ — report <b>reconciled</b>.",
  "तीनों match करते हैं — मतलब हमारी report पूरी तरह reconcile हो गई।",
  dash("B6", sheets=SHEETS1, hls=("C17", "G16", "B6")), cursor=[(0.2, "cell:B6", False)], sfx=[(0.3, "chime")])

# -------------------------------- Scene 12 : Conclusion
scene(12, "Scene 12 &bull; Conclusion")
DF = dash("A1", sheets=SHEETS1)
B("Toh doston, aaj humne zero se ek complete <b>Sales MIS Report</b> banayi.",
  "तो दोस्तों, आज हमने zero से एक complete Sales MIS Report बनाई।", DF, cursor=[(0.5, "cell:E2", False)])
B("Excel Table, structured formula, SUMIFS / COUNTIFS KPIs, region-product analysis, Pivot, charts aur slicers.",
  "Excel Table, structured reference formula, sum ifs और count ifs वाले KPIs, region और product analysis, pivot tables, charts और slicers।",
  DF, cursor=[(0.3, "cell:B6", False), (0.7, "cell:C14", False)])
B("Aur sabse zaroori — har number ko <b>validate</b> kiya. Yahi aapko professional <b>MIS executive</b> banata hai.",
  "और सबसे ज़रूरी — हमने हर number को validate किया। यही चीज़ आपको एक professional MIS executive बनाती है।",
  DF, cursor=[(0.5, "cell:G14", False)])
B("Video achhi lagi toh <b>Like</b>, <b>Subscribe</b> aur comment mein agla topic bataiye. Happy Excel-ing!",
  "अगर video अच्छी लगी, तो like कीजिए, channel subscribe कीजिए, और comment में बताइए कि अगली video किस topic पर चाहिए। मिलते हैं अगली video में — happy Excel-ing!",
  DF, cursor=[(0.6, "cell:E9", False)])

# ================================================================== cards
CARD_CSS = """*{margin:0;padding:0;box-sizing:border-box}html,body{width:1920px;height:1080px;overflow:hidden}
body{font-family:'Carlito','Calibri',sans-serif}
.hero{position:absolute;left:0;top:0;width:1920px;height:1080px;display:flex;flex-direction:column;align-items:center;
 justify-content:center;color:#fff;text-align:center}
.hero h1{font-size:92px;font-weight:900;letter-spacing:1px;text-shadow:0 4px 18px rgba(0,0,0,.4)}
.hero h2{font-size:38px;font-weight:600;margin-top:12px;color:#bfe3cc}
.hero .pill{margin-top:32px;background:rgba(255,255,255,.14);padding:13px 30px;border-radius:40px;font-size:25px}
.hero .sub{margin-top:24px;font-size:22px;opacity:.85}
.hero .list{margin-top:22px;font-size:28px;line-height:1.7;color:#dceee3}
.hero .cta{margin-top:36px;font-size:34px;background:#e8373a;padding:15px 44px;border-radius:50px;font-weight:800}"""
INTRO = ('<!doctype html><html><head><meta charset="utf-8"><style>' + CARD_CSS + '</style></head><body>'
         '<div class="hero" style="background:linear-gradient(135deg,#102a43,#1F3864 55%,#107C41)">'
         '<h1>SALES MIS REPORT</h1><h2>Microsoft Excel 365 &mdash; Complete Tutorial (Hinglish)</h2>'
         '<div class="pill">Excel Table &bull; SUMIFS / COUNTIFS &bull; Pivot &bull; Charts &bull; Slicers</div>'
         '<div class="sub">Animated recreation of Excel 365 &bull; AI voiceover &bull; numbers verified from the workbook</div>'
         '</div></body></html>')
OUTRO = ('<!doctype html><html><head><meta charset="utf-8"><style>' + CARD_CSS + '</style></head><body>'
         '<div class="hero" style="background:linear-gradient(135deg,#107C41,#1F3864 60%,#102a43)">'
         '<h1>Done! &#10003;</h1><div class="list">Excel Table &bull; Structured formulas &bull; KPIs &bull; Region / Product analysis<br>'
         'Pivot Tables &bull; Charts &bull; Slicers &mdash; every number <b>validated</b></div>'
         '<div class="cta">Like &bull; Subscribe &bull; Comment</div>'
         '<div class="sub">Happy Excel-ing! &mdash; see you in the next video</div></div></body></html>')
INTRO_DUR, OUTRO_DUR = 2.0, round(7.0 / SPEED, 2)   # title card: 2 seconds

# ================================================================== rendering helpers
def h(s):
    return hashlib.md5(s.encode()).hexdigest()[:16]


def chrome(args):
    return subprocess.run([CHROME, "--headless", "--no-sandbox", "--disable-gpu", "--hide-scrollbars",
                           "--force-device-scale-factor=1", *args], capture_output=True, text=True)


def render_png(html):
    k = h(html)
    p = f"{WORK}/png/{k}.png"
    if not os.path.exists(p):
        hp = f"{WORK}/html/{k}.html"
        open(hp, "w").write(html)
        chrome([f"--window-size=1920,{RENDER_H}", f"--screenshot={p}", f"file://{hp}"])
        assert os.path.exists(p), f"render failed {k}"
    return p


PROBE_JS = ("<script>document.addEventListener('DOMContentLoaded',function(){var o={};"
            "document.querySelectorAll('[data-id]').forEach(function(e){var r=e.getBoundingClientRect();"
            "if(r.width>0&&r.height>0)o[e.getAttribute('data-id')]=[r.left,r.top,r.width,r.height];});"
            "var p=document.createElement('pre');p.id='__R';p.textContent=JSON.stringify(o);document.body.appendChild(p);});</script>")
PROBE_CACHE_F = f"{WORK}/probe_cache.json"
PROBE_CACHE = json.load(open(PROBE_CACHE_F)) if os.path.exists(PROBE_CACHE_F) else {}


def probe(html):
    k = h(html)
    if k not in PROBE_CACHE:
        hp = f"{WORK}/html/probe_{k}.html"
        open(hp, "w").write(html.replace("</body>", PROBE_JS + "</body>"))
        out = chrome([f"--window-size=1920,{RENDER_H}", "--virtual-time-budget=1500", "--dump-dom", f"file://{hp}"]).stdout
        m = re.search(r'<pre id="__R">(.*?)</pre>', out, re.S)
        PROBE_CACHE[k] = json.loads(_html.unescape(m.group(1))) if m else {}
        json.dump(PROBE_CACHE, open(PROBE_CACHE_F, "w"))
    return PROBE_CACHE[k]


def resolve(target, frame, html):
    if isinstance(target, tuple):
        return target
    kind, ref = target.split(":", 1)
    if kind == "cell":
        x, y, w, hh = frame.v.rect(ref)
        r, c = ui.parse_ref(ref.split(":")[0])
        inf = frame.v.src.info(r, c)
        blank = frame.vkw.get("blank") or set()
        text = "" if (r, c) in blank else inf["text"]
        if ":" in ref or not text:
            return int(x + w / 2), int(y + hh / 2)
        hal = inf["h"] or ("right" if inf["num"] else "left")
        fx = 0.06 if hal == "right" else 0.90          # right-aligned value -> point left, else right
        return int(x + w * fx), int(y + hh * 0.55)
    rects = probe(html)
    if ref not in rects:
        raise KeyError(f"target {ref} not found in frame")
    x, y, w, hh = rects[ref]
    return int(x + w / 2), int(y + hh / 2)


# ================================================================== build
def main():
    # ---------- 1. voice
    print(f"[1/6] TTS for {len(BEATS)} beats (voice {VOICE}) ...", flush=True)
    voice = A.Voice(TTS_MODELS, f"{WORK}/tts", voice=VOICE, speed=VOICE_SPEED)
    for i, b in enumerate(BEATS):
        b["wav"], b["speech"] = voice.say(pronunciation.apply(b["tts"]))

    # ---------- 2. timeline
    t = INTRO_DUR
    for b in BEATS:
        sp = b["speech"]
        fr = [f for f, _ in b["frames"]] + [f for f, _, _ in b["cursor"]] + [f for f, _ in b["sfx"]]
        last = max(fr + [0.0])
        b["dur"] = round(max(LEAD + sp + TAIL, LEAD + last * sp + 0.9 / SPEED) + b["extra"], 3)
        b["start"] = t
        t += b["dur"]
    total = t + OUTRO_DUR
    print(f"      narration {sum(b['speech'] for b in BEATS):.1f}s, video {total:.1f}s = {total/60:.2f} min", flush=True)

    # ---------- 3. frames + cursor
    print("[2/6] rendering frames + resolving cursor targets ...", flush=True)
    cur = (960, 640)
    plans = []
    sfx_events = [(0.25, "whoosh")]
    prev_scene = None
    prev_dialog = False
    for bi, b in enumerate(BEATS):
        sp = b["speech"]
        T = lambda f: LEAD + f * sp
        htmls = [(T(f) if f > 0 else 0.0, fr, fr.html(b["caption"], b["badge"])) for f, fr in b["frames"]]
        htmls.sort(key=lambda x: x[0])
        pngs = [(tt, render_png(hh)) for tt, _, hh in htmls]
        for tt, _, hh in htmls:
            dlg = 'class="dlg"' in hh
            if dlg and not prev_dialog:
                pass
            prev_dialog = dlg

        def frame_at(tt):
            cand = [x for x in htmls if x[0] <= tt + 1e-6]
            return cand[-1] if cand else htmls[0]

        way, clicks = [(0.0, cur[0], cur[1])], []
        for f, target, click in sorted(b["cursor"], key=lambda x: x[0]):
            ta = T(f)
            _, fr, hh = frame_at(ta - 0.01)
            x, y = resolve(target, fr, hh)
            way.append((ta, x, y))
            if click:
                clicks.append((ta, x, y))
                sfx_events.append((b["start"] + ta, "click"))
        cur = (way[-1][1], way[-1][2])
        for f, kind in b["sfx"]:
            sfx_events.append((b["start"] + T(f), kind))
        if b["scene"] != prev_scene:
            sfx_events.append((b["start"] + 0.02, "whoosh"))
            prev_scene = b["scene"]
        plans.append(dict(pngs=pngs, way=way, clicks=clicks, dur=b["dur"]))
        if (bi + 1) % 10 == 0:
            print(f"      {bi+1}/{len(BEATS)} beats", flush=True)

    # ---------- 4. clips
    print("[3/6] encoding clips ...", flush=True)
    cursor_png, rings = make_cursor_assets()
    clips = [card_clip(INTRO, INTRO_DUR, "intro")]
    for i, p in enumerate(plans):
        clips.append(beat_clip(i, p, cursor_png, rings))
    clips.append(card_clip(OUTRO, OUTRO_DUR, "outro"))
    lst = f"{WORK}/clips/list.txt"
    open(lst, "w").write("".join(f"file '{c}'\n" for c in clips))
    silent = f"{WORK}/video_silent.mp4"
    subprocess.run(["ffmpeg", "-y", "-v", "error", "-f", "concat", "-safe", "0", "-i", lst, "-c", "copy", silent], check=True)

    # ---------- 5. audio
    print("[4/6] building voice / sfx / music tracks ...", flush=True)
    vt = np.zeros(int(total * A.TTS_SR) + A.TTS_SR, np.float32)
    for b in BEATS:
        a, _ = sf.read(b["wav"], dtype="float32")
        i = int((b["start"] + LEAD) * A.TTS_SR)
        vt[i:i + len(a)] += a
    vt = vt[: int(total * A.TTS_SR)]
    sf.write(f"{WORK}/audio/voice.wav", vt, A.TTS_SR)
    sfx_events.append((total - OUTRO_DUR + 0.1, "whoosh"))
    sf.write(f"{WORK}/audio/sfx.wav", A.render_sfx(sfx_events, total), A.SR)
    sf.write(f"{WORK}/audio/music.wav", A.render_music(total), A.SR)
    mix = f"{WORK}/audio/mix.wav"
    fc = ("[0:a]aresample=48000:resampler=soxr,pan=stereo|c0=c0|c1=c0,asplit=2[v][key];"
          "[1:a]pan=stereo|c0=c0|c1=c0,volume=0.55[s];"
          # bed ~20 dB under the voice; lifted only on the title/end cards where nobody speaks
          f"[2:a]volume='0.10*if(lt(t,{INTRO_DUR}),4.0,if(gt(t,{total - OUTRO_DUR:.2f}),4.0,1))':eval=frame[m];"
          "[m][key]sidechaincompress=threshold=0.01:ratio=10:attack=15:release=700:makeup=1[md];"
          "[v][s][md]amix=inputs=3:normalize=0:duration=first[a]")
    subprocess.run(["ffmpeg", "-y", "-v", "error", "-i", f"{WORK}/audio/voice.wav", "-i", f"{WORK}/audio/sfx.wav",
                    "-i", f"{WORK}/audio/music.wav", "-filter_complex", fc.replace("[md];", "[md0];[md0]asplit=2[md][mdq];"),
                    "-map", "[a]", "-ar", "48000", mix, "-map", "[mdq]", "-ar", "48000", f"{WORK}/audio/music_ducked.wav"],
                   check=True)
    # Fixed-gain normalisation to -14 LUFS + peak limiter. (loudnorm can silently fall back
    # to dynamic mode, which pumps the music-only intro up to speech level.)
    meas = subprocess.run(["ffmpeg", "-hide_banner", "-nostats", "-i", mix, "-af", "ebur128", "-f", "null", "-"],
                          capture_output=True, text=True).stderr
    integ = float(re.findall(r"I:\s+(-?[0-9.]+) LUFS", meas)[-1])
    gain_db = -14.0 - integ
    ln = f"volume={gain_db:.2f}dB,alimiter=limit=0.70:attack=5:release=60:level=false"
    print(f"      mix {integ:.1f} LUFS -> gain {gain_db:+.2f} dB", flush=True)

    # ---------- 6. mux + subtitles
    print("[5/6] muxing final MP4 ...", flush=True)
    subprocess.run(["ffmpeg", "-y", "-v", "error", "-i", silent, "-i", mix, "-map", "0:v", "-map", "1:a",
                    "-c:v", "copy", "-af", ln, "-c:a", "aac", "-b:a", "192k", "-ar", "48000",
                    "-movflags", "+faststart", "-shortest", OUT], check=True)
    print("[6/6] writing synced subtitles ...", flush=True)
    write_srt(BEATS)
    json.dump([dict(i=i + 1, scene=b["scene"], start=round(b["start"], 3), dur=b["dur"], speech=round(b["speech"], 3),
                    wav=b["wav"], caption=re.sub(r"<[^>]+>", "", b["caption"]), tts=b["tts"])
               for i, b in enumerate(BEATS)], open(f"{WORK}/timeline.json", "w"), ensure_ascii=False, indent=1)
    d = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=noprint_wrappers=1:nokey=1", OUT],
                       capture_output=True, text=True).stdout.strip()
    print(f"DONE: {OUT}  ({float(d or 0):.1f}s = {float(d or 0)/60:.2f} min)", flush=True)


# ---------------------------------------------------------------- clip helpers
def make_cursor_assets():
    """Cursor with a yellow highlight halo (like the reference recording) + 3 click-ring frames.
    Drawn with PIL at 4x and downsampled: headless Chromium paints nothing in tiny windows."""
    from PIL import Image, ImageDraw
    S4 = 4
    def canvas():
        return Image.new("RGBA", (96 * S4, 96 * S4), (0, 0, 0, 0))
    cur = f"{WORK}/assets/cursor_v3.png"
    im = canvas()
    d = ImageDraw.Draw(im)
    c = 48 * S4
    d.ellipse([c - 23 * S4, c - 23 * S4, c + 23 * S4, c + 23 * S4], fill=(255, 214, 0, 80))
    arrow = [(48, 48), (48, 82), (56, 74), (62, 88), (68, 85), (62, 71), (74, 71)]
    pts = [(x * S4, y * S4) for x, y in arrow]
    d.polygon(pts, fill=(255, 255, 255, 255))
    d.line(pts + [pts[0]], fill=(0, 0, 0, 255), width=int(2.2 * S4), joint="curve")
    im.resize((96, 96), Image.LANCZOS).save(cur)
    rings = []
    for i, r in enumerate((16, 26, 36)):
        pth = f"{WORK}/assets/ring{i}_v2.png"
        im = canvas()
        d = ImageDraw.Draw(im)
        a = int(255 * (0.9 - i * 0.25))
        d.ellipse([c - r * S4, c - r * S4, c + r * S4, c + r * S4], outline=(232, 140, 0, a), width=4 * S4)
        im.resize((96, 96), Image.LANCZOS).save(pth)
        rings.append(pth)
    return cur, rings


def _pw(way, idx):
    """Piecewise smoothstep expression for coordinate idx (1=x, 2=y) over waypoints."""
    expr = str(way[-1][idx])
    for k in range(len(way) - 1, 0, -1):
        t0, t1 = way[k - 1][0], way[k][0]
        a, b = way[k - 1][idx], way[k][idx]
        ms = max(t0, t1 - min(0.85 / SPEED, max(0.25, t1 - t0)))
        span = max(t1 - ms, 1e-3)
        p = f"((t-{ms:.3f})/{span:.3f})"
        moving = f"{a}+({b - a})*{p}*{p}*(3-2*{p})"
        expr = f"if(lt(t,{ms:.3f}),{a},if(lt(t,{t1:.3f}),{moving},{expr}))"
    return expr


def beat_clip(i, p, cursor_png, rings):
    key = h(json.dumps([p["pngs"], p["way"], p["clicks"], p["dur"]]))
    out = f"{WORK}/clips/b{i:02d}_{key}.mp4"
    if os.path.exists(out):
        return out
    lst = f"{WORK}/clips/b{i:02d}.txt"
    pngs = p["pngs"]
    lines = []
    for j, (tt, png) in enumerate(pngs):
        nxt = pngs[j + 1][0] if j + 1 < len(pngs) else p["dur"]
        lines.append(f"file '{png}'\nduration {max(nxt - tt, 1 / FPS):.4f}\n")
    lines.append(f"file '{pngs[-1][1]}'\n")
    open(lst, "w").write("".join(lines))
    xe, ye = _pw(p["way"], 1), _pw(p["way"], 2)
    inputs = ["-f", "concat", "-safe", "0", "-i", lst, "-i", cursor_png] + sum((["-i", r] for r in rings), [])
    fc = [f"[0:v]crop=1920:1080:0:0,fps={FPS},format=yuva420p[b0]"]
    last = "b0"
    n = 0
    for (tc, x, y) in p["clicks"]:
        for k, r in enumerate(rings):
            a, bnd = tc + k * 0.08, tc + (k + 1) * 0.08 + 0.02
            n += 1
            fc.append(f"[{last}][{2+k}:v]overlay=x={x-48}:y={y-48}:enable='between(t,{a:.3f},{bnd:.3f})'[r{n}]")
            last = f"r{n}"
    fc.append(f"[{last}][1:v]overlay=x='{xe}-48':y='{ye}-48':eval=frame,format=yuv420p[v]")
    subprocess.run(["ffmpeg", "-y", "-v", "error", *inputs, "-filter_complex", ";".join(fc), "-map", "[v]",
                    "-t", f"{p['dur']:.3f}", "-r", str(FPS), "-c:v", "libx264", "-crf", "19", "-preset", "veryfast",
                    "-pix_fmt", "yuv420p", out], check=True)
    return out


def card_clip(html, dur, name):
    png = render_png(html)
    out = f"{WORK}/clips/{name}_{h(html)}_{dur}.mp4"
    if not os.path.exists(out):
        subprocess.run(["ffmpeg", "-y", "-v", "error", "-loop", "1", "-t", str(dur), "-i", png, "-vf",
                        f"crop=1920:1080:0:0,fps={FPS},fade=t=in:st=0:d=0.5,fade=t=out:st={dur-0.6}:d=0.6,format=yuv420p",
                        "-r", str(FPS), "-c:v", "libx264", "-crf", "19", "-preset", "veryfast", "-pix_fmt", "yuv420p", out],
                       check=True)
    return out


def write_srt(beats):
    def ts(s):
        ms = int(round(s * 1000))
        return f"{ms//3600000:02d}:{ms//60000%60:02d}:{ms//1000%60:02d},{ms%1000:03d}"
    out = []
    for i, b in enumerate(beats, 1):
        s = b["start"] + LEAD
        e = s + b["speech"]
        txt = re.sub(r"<[^>]+>", "", b["caption"]).replace("&amp;", "&")
        out.append(f"{i}\n{ts(s)} --> {ts(e)}\n{txt}\n")
    open(SRT, "w").write("\n".join(out))


if __name__ == "__main__":
    main()
